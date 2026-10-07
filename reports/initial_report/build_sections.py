from pathlib import Path
import csv,re,hashlib,json
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parent;root=P.parents[1]
(P/'figures').mkdir(exist_ok=True)
rows=[r for r in csv.DictReader((root/'data/current/category_summary.csv').open()) if r['priority']=='YES']
rows.sort(key=lambda r:int(r['annotation_records']))
assert len(rows)==32 and sum(int(r['annotation_records']) for r in rows)==2290
plt.rcParams.update({'font.size':11,'font.family':'DejaVu Sans'})
fig,ax=plt.subplots(figsize=(9,10.8))
ax.barh([r['category'] for r in rows],[int(r['annotation_records']) for r in rows],color='#16756D')
for i,r in enumerate(rows):ax.text(int(r['annotation_records'])+3,i,r['annotation_records'],va='center',fontsize=10)
ax.set_xlabel('Number of annotation records');ax.set_title('Fall 2026: annotations across 32 target categories',pad=16)
ax.set_xlim(0,350);ax.spines[['top','right']].set_visible(False);ax.set_axisbelow(True);ax.xaxis.grid(alpha=.15)
fig.tight_layout();fig.savefig(P/'figures/category_counts.png',dpi=220);fig.savefig(P/'figures/category_counts.pdf');plt.close(fig)
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7)
for st in ['Heading 1','Heading 2']:doc.styles[st].font.color.rgb=RGBColor.from_string('16756D')
lines=(P/'INTRO_BACKGROUND_DATA.md').read_text().splitlines();i=0
while i<len(lines):
 line=lines[i]
 if not line.strip():i+=1;continue
 if line.startswith('|'):
  tab=[]
  while i<len(lines) and lines[i].startswith('|'):
   values=[x.strip() for x in lines[i].strip('|').split('|')]
   if not all(re.fullmatch(r'[:\- ]+',x) for x in values):tab.append(values)
   i+=1
  t=doc.add_table(rows=0,cols=len(tab[0]));t.style='Light Shading Accent 1'
  for values in tab:
   for c,v in zip(t.add_row().cells,values):c.text=v
  continue
 if line.startswith('!['):
  doc.add_page_break();doc.add_picture(str(P/'figures/category_counts.png'),width=Inches(6.4))
 elif line.startswith('# '):doc.add_heading(line[2:],0)
 elif line.startswith('## '):doc.add_heading(line[3:],1)
 else:
  p=doc.add_paragraph()
  for j,part in enumerate(line.split('**')):p.add_run(part).bold=bool(j%2)
 i+=1
foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('Draft sections | ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
doc.save(P/'INTRO_BACKGROUND_DATA.docx')
rubric=P/'Report Rubrics.pdf'
(P/'source_receipt.json').write_text(json.dumps({'original_filename':rubric.name,'received':'2026-09-27','sha256':hashlib.sha256(rubric.read_bytes()).hexdigest(),'use':'Requirements for initial report; sections drafted separately for insertion into existing team report.'},indent=2)+'\n')
print('Created editable Word sections and category count figure.')
