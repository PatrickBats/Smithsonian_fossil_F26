from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
p=Presentation();p.slide_width=Inches(10);p.slide_height=Inches(5.625)
C={'bg':'F8F5EF','ink':'202D48','accent':'93451F','light':'EEE3D6','muted':'596477','white':'FFFFFF'}
def tx(s,x,y,w,h,t,size=16,color='ink',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line;q.font.name='Lato';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(C[color]);q.space_after=Pt(4)
 return sh

def rect(s,x,y,w,h,c):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C[c]);sh.line.fill.background();return sh

def slide(n,tag,title):
 s=p.slides.add_slide(p.slide_layouts[6]);rect(s,0,0,10,5.625,'bg');rect(s,.4,.29,.25,.04,'accent');tx(s,.77,.2,8.7,.2,f'{n:02d} / {tag}',8,'accent',True)
 tx(s,.4,.65,9.2,.66,title,27,bold=True);rect(s,.4,5.26,9.2,.009,'light');tx(s,.4,5.37,8,.16,'SMITHSONIAN × RICE D2K / SPONSOR UPDATE',7,'muted');tx(s,9.35,5.34,.3,.2,str(n),8,'muted');return s

def pic(s,num,x,y,w):
 file=ROOT/f'reports/initial_report_complete_2026-09-27/figures/specimen_sources/shard_0_target_{num}_crop.png'
 with Image.open(file) as im:a,b=im.size
 scale=min(w/a,w/b);ww=a*scale;hh=b*scale
 s.shapes.add_picture(str(file),Inches(x+(w-ww)/2),Inches(y+(w-hh)/2),width=Inches(ww),height=Inches(hh))

s=slide(1,'PROGRESS','This week: presentation and data preparation')
rect(s,.4,1.62,4.2,2.69,'white');tx(s,.62,1.88,3.76,.36,'Main focus this week',19,bold=True)
tx(s,.62,2.48,3.73,1.1,'Prepared and presented our\ninitial class slides and report.',20)
tx(s,.62,3.72,3.73,.34,'Clarified the scope and evaluation plan.',12,'muted')
tx(s,5.05,1.77,4.4,.4,'Initial specimen crops prepared',19,bold=True)
tx(s,5.1,2.22,4.2,.18,'Example specimens',10,'muted')
for i,n in enumerate(['008','016','027']):pic(s,n,5.1+i*1.47,2.41,1.3)
tx(s,5.1,3.94,4.2,.42,'2,283 single-plane crops · split under review',13,'muted')
tx(s,.45,4.79,9.1,.3,'Full evaluation and classifier training are still ahead.',17,'accent',True)

s=slide(2,'PRELIMINARY DETECTION','Previous detectors recover many known specimens')
tx(s,.45,1.38,9.1,.38,'Small diagnostic: 64 annotated targets across 29 microscope slides',17,'muted')
for y,label,n in [(2.12,'Best-plane checkpoint',54),(3.14,'Focus-stack checkpoint',51)]:
 tx(s,.45,y,3.1,.4,label,18,bold=True)
 rect(s,3.67,y+.025,4.25,.38,'light');rect(s,3.67,y+.025,4.25*n/64,.38,'accent')
 tx(s,8.17,y-.07,1.5,.52,f'{n}/64',27,bold=True)
 tx(s,8.17,y+.45,1.4,.29,f'{n/64:.1%} recovered',11,'muted')
tx(s,.45,4.29,9.1,.39,'Encouraging for reuse; broader evaluation is still needed.',19,'accent',True)
tx(s,.45,4.91,9.1,.21,'Known-target recovery in centered crops—not whole-slide or classification accuracy.',11,'muted')

s=slide(3,'NEXT STEPS','Finalize the training, validation, and test sets')
for x,n,title,body in [(.45,'01','Review the examples','Check crop quality\nand uncertain labels.'),(3.6,'02','Finalize the split','Group related slides;\ncheck category coverage.'),(6.75,'03','Start the baseline','Train the first classifier\nand evaluate its errors.')]:
 rect(s,x,1.69,2.8,2.57,'white');tx(s,x+.18,1.91,2.44,.28,n,12,'accent',True);tx(s,x+.18,2.43,2.44,.71,title,20,bold=True);tx(s,x+.18,3.35,2.44,.68,body,14,'muted')
tx(s,.45,4.66,9.1,.39,'Also: resolve shared storage for the full multifocal crop export.',15,'accent',True)
notes=[
'This week, much of our effort went into preparing and presenting our initial class slides and report. That helped us clarify the scope and how we want to evaluate the models. We also made progress with data preparation. The latest team update reports 2,283 single-plane specimen crops and a draft split that we’re reviewing. We haven’t completed a full evaluation or trained the category classifiers yet.',
'Our preliminary check used two frozen detectors from the previous work on 64 annotated targets across 29 slides. The best-plane checkpoint recovered 54 targets, and the focus-stack checkpoint recovered 51. This is an encouraging sign that we can reuse the detection work. However, these were small regions centered on known specimens, so the numbers are not whole-slide accuracy or classification accuracy. The checkpoints have different weights, so this also doesn’t isolate which image-preparation method is better.',
'Our next priority is to finalize the training, validation, and test sets. We’ll review crop quality and label issues, keep related microscope slides together, and check that the category coverage supports a useful evaluation. Once that protocol is agreed, we can train the first classification baseline. We also need to resolve shared storage before completing the multifocal crop export.'
]
sources=[
'Team progress: PR #1, commit be4634b50e8f5f4404216f8e5ba13fc2e31410dd. Crop completion is reported by the contributor; the review independently reproduced the 2,283-record candidate count. Crops are single-plane, not focus-selected or fully multifocal. Split remains a review draft.',
'Source: reports/detection_pilot_2026-09-24/README.md. Confidence >= 0.50 and box IoU >= 0.50. Frozen YOLO checkpoints, not RF-DETR or classifier results. 50 targets overlap historical training-slide membership, 9 historical test, 5 unresolved; checkpoint exposure not fully verified. These results cannot establish independent generalization.',
'PR #1: draft 1,605 train / 347 validation / 331 test; one category has no validation support and two test examples. These are not approved final assignments. Storage quotas were reported full October 1; not rechecked for this slide. No new jobs authorized by this presentation.'
]
seconds=[40,45,35]
md=['# Sponsor update — speaker notes\n','Three slides · approximately two minutes. Source context is not spoken.\n']
for i,(s,n,src,sec) in enumerate(zip(p.slides,notes,sources,seconds)):
 s.notes_slide.notes_text_frame.text=f'TARGET: {sec} seconds\n\nSAY:\n{n}\n\nSOURCE CONTEXT — NOT SPOKEN:\n{src}'
 md.append(f'## Slide {i+1} — {sec} seconds\n\n{n}\n')
p.save(OUT/'Smithsonian_sponsor_update.pptx')
(OUT/'speaker_notes.md').write_text('\n'.join(md))
(OUT/'timing.json').write_text(json.dumps({'seconds':seconds,'total_seconds':sum(seconds),'spoken_words':sum(len(n.split()) for n in notes)},indent=2)+'\n')
print('3 slides;',sum(len(n.split()) for n in notes),'spoken words; 120 seconds planned')
