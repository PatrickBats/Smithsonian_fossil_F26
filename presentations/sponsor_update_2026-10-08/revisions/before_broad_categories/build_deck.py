from pathlib import Path
import json, textwrap, hashlib
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image

OUT=Path(__file__).resolve().parent;BASE=OUT.parents[2]
C={'bg':'F8F5EF','ink':'202D48','accent':'93451F','muted':'596477','line':'E6DFD4','white':'FFFFFF','teal':'20786E','pale':'E4EFEB','gray':'9AA5B4'}
pdfmetrics.registerFont(TTFont('Lato','/usr/share/fonts/truetype/lato/Lato-Regular.ttf'))
pdfmetrics.registerFont(TTFont('LatoB','/usr/share/fonts/truetype/lato/Lato-Bold.ttf'))
p=Presentation();p.slide_width=Inches(960/72);p.slide_height=Inches(540/72)
c=canvas.Canvas(str(OUT/'Smithsonian_sponsor_update.pdf'),pagesize=(960,540));c.setTitle('Smithsonian: classification progress')
notes=[];times=[25,0,45,50,55,45];titles=[]
def col(k):return '#'+C.get(k,k)
def rect(s,x,y,w,h,k):
 z=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));z.fill.solid();z.fill.fore_color.rgb=RGBColor.from_string(C.get(k,k));z.line.fill.background()
 c.setFillColor(col(k));c.rect(x,540-y-h,w,h,fill=1,stroke=0)
def txt(s,x,y,w,h,t,size=18,k='ink',bold=False):
 z=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));f=z.text_frame;f.word_wrap=False
 f.margin_top=f.margin_bottom=f.margin_left=f.margin_right=0
 for i,line in enumerate(t.split('\n')):
  q=f.paragraphs[0] if i==0 else f.add_paragraph();q.text=line;q.font.name='Lato';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(C.get(k,k));q.space_before=q.space_after=Pt(0);q.line_spacing=1.12
  assert pdfmetrics.stringWidth(line,'LatoB' if bold else 'Lato',size)<=w+4,(line,w)
  c.setFont('LatoB' if bold else 'Lato',size);c.setFillColor(col(k));c.drawString(x,540-y-size*.9-i*size*1.12,line)
 return z
def pic(s,name,x,y,w,h):
 path=OUT/'assets'/name
 with Image.open(path) as im:iw,ih=im.size
 ratio=min(w/iw,h/ih);ww,hh=iw*ratio,ih*ratio;x+=(w-ww)/2;y+=(h-hh)/2
 s.shapes.add_picture(str(path),Pt(x),Pt(y),Pt(ww),Pt(hh));c.drawImage(str(path),x,540-y-hh,width=ww,height=hh)
def start(tag,title):
 n=len(p.slides)+1;s=p.slides.add_slide(p.slide_layouts[6]);rect(s,0,0,960,540,'bg');rect(s,42,30,26,4,'accent')
 txt(s,80,23,830,22,f'{n:02d} / {tag}',11,'accent',True);txt(s,42,65,880,88,title,34,bold=True)
 rect(s,42,505,876,1,'line');txt(s,42,518,800,14,'SMITHSONIAN × RICE D2K  ·  CLASSIFICATION UPDATE',9,'muted');txt(s,900,517,20,16,str(n),10,'muted')
 titles.append(title.replace('\n',' '));return s
def finish(s,note):
 notes.append(note);s.notes_slide.notes_text_frame.text=f'TARGET: {times[len(notes)-1]} seconds\n\n{note}';c.showPage()

s=start('PROJECT UPDATE','Two approaches to identifying fossil pollen')
txt(s,44,127,860,32,'RF-DETR method first · classifier backbones second',21,'muted')
for y,number,title,body in [(193,'01','RF-DETR','Find specimens and identify categories.'),(316,'02','Classifier backbones','Identify categories from specimen crops.')]:
 rect(s,44,y,374,100,'white');txt(s,59,y+14,40,29,number,18,'teal',True)
 txt(s,105,y+12,300,34,title,24,bold=True);txt(s,59,y+62,344,25,body,16,'muted')
for i,(name,label) in enumerate([('hero_0.png','Bisaccate'),('hero_1.png','Bombacacidites sp.'),('hero_2.png','Arecipites sp.')]):
 x=454+i*155;rect(s,x,214,145,190,'white');pic(s,name,x+6,221,133,143);txt(s,x+7,374,139,26,label,11,'muted')
txt(s,44,450,868,29,'Classification dataset: 2,283 specimens · 32 priority categories',18,'teal',True)
finish(s,"We’ve organized this update into two parts. First is the RF-DETR method, which aims to locate specimens and identify their categories together. Then we’ll cover the classifier backbones, which identify categories from prepared specimen crops. That classification experiment uses 2,283 specimens across 32 priority categories. The images here are real examples with expert labels.")

s=start('SECTION 01 · RF-DETR','RF-DETR: detection and category identification')
rect(s,44,187,874,222,'white')
txt(s,79,267,800,48,'RF-DETR team update',32,'muted',True)
txt(s,79,329,800,30,'Placeholder · content to be added',20,'muted')
finish(s,"PLACEHOLDER — not a spoken script. RF-DETR team to add its method, progress and results. Allow approximately 60 seconds if the full presentation should stay near five minutes. No RF-DETR results are supplied or implied here.")

s=start('SECTION 02 · CLASSIFIER BACKBONES / FOCUS','A clearer view helps the model identify pollen')
for x,name,label in [(44,'focus_before.png','Original focal plane'),(302,'focus_after.png','Selected sharper plane')]:
 txt(s,x,158,248,30,label,18,bold=True);pic(s,name,x,198,230,230)
txt(s,44,451,510,30,'Same specimen and crop · expert label: Quercus sp.',14,'muted')
rect(s,577,166,341,276,'pale');txt(s,601,190,295,32,'Unseen-slide accuracy',19,'teal',True)
txt(s,601,244,292,68,'32% → 38%',43,'teal',True)
txt(s,601,331,295,82,'Same Swin model settings.\nOnly the selected focal\nplane changed.',19)
finish(s,"The microscope scans show each specimen at several focal depths. On the left is the original view; on the right is a sharper view of the same specimen. We used a sharpness score to choose one plane for each crop. With the same Swin training settings, validation accuracy on unfamiliar slides increased from about 32 to 38 percent. This example illustrates the focus change; it doesn’t mean this particular specimen caused that improvement. We’d also like to know when identification requires features visible across several depths, rather than just one sharp image.")

s=start('CLASSIFIER BACKBONES / EVALUATION','Two questions, two ways to check performance')
for x,heading,sub in [(44,'Unseen slides','Can it work on a different slide?'),(492,'Shared slides','Can it identify more grains on familiar slides?')]:
 rect(s,x,164,424,283,'white');txt(s,x+20,183,390,34,heading,25,'teal',True);txt(s,x+20,226,388,25,sub,15,'muted')
 txt(s,x+20,280,160,25,'LEARN FROM',12,'muted',True);txt(s,x+230,280,170,25,'CHECK ON',12,'muted',True)
for x,labels in [(64,['Slide A','Slide B']),(274,['Slide C','Slide D']),(512,['A · grain 1','B · grain 1']),(722,['A · grain 2','B · grain 2'])]:
 for j,label in enumerate(labels):
  rect(s,x,315+j*48,168,36,'pale' if x in [274,722] else 'line');txt(s,x+12,323+j*48,147,25,label,16)
txt(s,44,467,874,25,'Different specimens in every set. Final test slides remain reserved.',18,'accent',True)
finish(s,"We checked performance in two situations. In the unseen-slide experiment, all specimens from a slide stay together, so the model is checked on slides it hasn’t learned from. In the shared-slide experiment, it learns from some grains and is checked on other grains from the same slides. We never reuse the exact same specimen in both sets. Shared slides can have familiar staining and imaging conditions, so that is an easier situation. Both experiments have the same number of training and validation examples, but the actual examples differ. A separate final test set remains unused.")

s=start('CLASSIFIER BACKBONES / RESULTS','Adapting DINO gives our strongest results so far')
txt(s,42,133,880,26,'Validation accuracy · identifying categories in prepared specimen crops',17,'muted')
txt(s,315,186,245,30,'Unseen slides',23,bold=True);txt(s,646,186,260,30,'Shared slides',23,bold=True)
models=[('Swin',38.0403458213,66.2824207493),('DINO · frozen',29.9711815562,59.0778097983),('DINO · adapted',43.5158501441,70.0288184438)]
for i,(label,left,right) in enumerate(models):
 y=249+i*66;txt(s,43,y+2,249,35,label,23,bold=i==2)
 for x,v in [(315,left),(646,right)]:
  rect(s,x,y,210,28,'line');rect(s,x,y,210*v/100,28,'teal' if i==2 else 'gray');txt(s,x+220,y-2,66,36,f'{v:.0f}%',25,'teal' if i==2 else 'ink',True)
txt(s,44,454,870,25,'Adapted = trained the final two DINO blocks and its category classifier.',16,'muted')
txt(s,44,481,880,17,'Single-run validation comparisons; final test untouched. These are not whole-slide detection results.',11,'muted')
finish(s,"Here are the classification results. Swin reached 38 percent on unseen slides and 66 percent with shared slides. DINO initially did worse when we kept its image-processing backbone fixed and trained only a small classifier. We then allowed the final part of DINO to adapt to our specimens. That reached about 44 percent on unseen slides and 70 percent with shared slides. Its category-averaged F1 score also improved. These are promising initial results, but they are single-run validation comparisons, not final test results. They measure identification of prepared crops, rather than the complete process of finding and classifying everything on a slide.")

s=start('CLASSIFIER BACKBONES / ERRORS','Where identification still struggles')
txt(s,44,130,874,30,'Adapted DINO · unseen-slide validation',20,'muted')
txt(s,44,185,425,24,'EXPERT CATEGORY',12,'muted',True)
txt(s,483,185,230,24,'CORRECTLY IDENTIFIED',12,'muted',True)
txt(s,790,185,128,24,'CORRECT / TOTAL',12,'muted',True)
metrics_path=BASE/'dino_finetune_2026-10-07/separate_slides/backbone_finetune/best_validation_metrics.json'
metrics=json.loads(metrics_path.read_text())
worst=sorted([v for v in metrics['per_class'] if v['support']>=5],key=lambda v:(v['recall'],-v['support']))[:6]
for i,v in enumerate(worst):
 y=222+i*35
 txt(s,44,y,421,31,v['category'],20,bold=i==0)
 rect(s,483,y+2,224,18,'line')
 if v['recall']>0:rect(s,483,y+2,224*v['recall'],18,'accent')
 txt(s,721,y-1,62,30,f"{v['recall']:.0%}",20,'accent',True)
 correct=round(v['recall']*v['support']);txt(s,809,y-1,110,30,f"{correct} / {v['support']}",20,bold=True)
txt(s,44,449,870,23,'18 of 21 Momipites wyomingensis were predicted as Momipites sp.',17,'accent',True)
txt(s,44,479,875,18,'Lowest category recall among categories with ≥5 validation examples; small samples remain uncertain.',12,'muted')
finish(s,"This shows the categories the adapted DINO model struggled with most on unfamiliar slides. Momipites wyomingensis was the clearest problem: none of its 21 validation specimens were identified correctly, and 18 were called Momipites instead. Polyatriopollenites, Bombacacidites, Ulmipollenites, Psilatricolpites and Arecipites also had low recovery. The counts on the right show how many examples each result is based on. We’ve limited this ranking to categories with at least five examples, since a result based on one or two specimens would be especially unreliable. These are the main categories to investigate next.")

p.save(OUT/'Smithsonian_sponsor_update.pptx');c.save()
(OUT/'speaker_notes.md').write_text('# Sponsor update speaker notes\n\nPrepared content: approximately 3:40. RF-DETR slide is an unfilled placeholder; allow about 60 seconds for that team’s update to target 4:40 total. No meeting date assumed.\n\n'+'\n\n'.join(f'## {i+1}. {t} — {times[i]} seconds\n\n{n}' for i,(t,n) in enumerate(zip(titles,notes))))
(OUT/'timing.json').write_text(json.dumps({'seconds':sum(times),'spoken_words':sum(len(n.split()) for n in notes),'slides':[{'slide':i+1,'title':t,'seconds':times[i],'words':len(notes[i].split())} for i,t in enumerate(titles)]},indent=2)+'\n')
print('Generated',len(p.slides),'slides;',sum(len(n.split()) for n in notes),'spoken words')
