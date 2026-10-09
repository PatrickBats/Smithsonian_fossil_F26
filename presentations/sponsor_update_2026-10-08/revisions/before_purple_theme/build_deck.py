from pathlib import Path
import json, textwrap, hashlib
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image

OUT=Path(__file__).resolve().parent;BASE=OUT.parents[2]
C={'bg':'F8F5EF','ink':'202D48','accent':'93451F','muted':'596477','line':'E6DFD4','white':'FFFFFF','teal':'20786E','pale':'E4EFEB','gray':'9AA5B4'}
pdfmetrics.registerFont(TTFont('Lato','/usr/share/fonts/truetype/lato/Lato-Regular.ttf'))
pdfmetrics.registerFont(TTFont('LatoB','/usr/share/fonts/truetype/lato/Lato-Bold.ttf'))
p=Presentation();p.slide_width=Inches(960/72);p.slide_height=Inches(540/72)
c=canvas.Canvas(str(OUT/'Smithsonian_sponsor_update.pdf'),pagesize=(960,540));c.setTitle('Smithsonian: classification progress')
notes=[];times=[25,0,45,50,55,45];titles=[]
def col(k):return '#'+C.get(k,k)
def rect(s,x,y,w,h,k):
 z=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));z.fill.solid();z.fill.fore_color.rgb=RGBColor.from_string(C.get(k,k));z.line.fill.background()
 c.setFillColor(col(k));c.rect(x,540-y-h,w,h,fill=1,stroke=0)
def txt(s,x,y,w,h,t,size=18,k='ink',bold=False):
 z=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));f=z.text_frame;f.word_wrap=False
 f.margin_top=f.margin_bottom=f.margin_left=f.margin_right=0
 for i,line in enumerate(t.split('\n')):
  q=f.paragraphs[0] if i==0 else f.add_paragraph();q.text=line;q.font.name='Lato';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(C.get(k,k));q.space_before=q.space_after=Pt(0);q.line_spacing=1.12
  assert pdfmetrics.stringWidth(line,'LatoB' if bold else 'Lato',size)<=w+4,(line,w)
  c.setFont('LatoB' if bold else 'Lato',size);c.setFillColor(col(k));c.drawString(x,540-y-size*.9-i*size*1.12,line)
 return z
def pic(s,name,x,y,w,h):
 path=OUT/'assets'/name
 with Image.open(path) as im:iw,ih=im.size
 ratio=min(w/iw,h/ih);ww,hh=iw*ratio,ih*ratio;x+=(w-ww)/2;y+=(h-hh)/2
 s.shapes.add_picture(str(path),Pt(x),Pt(y),Pt(ww),Pt(hh));c.drawImage(str(path),x,540-y-hh,width=ww,height=hh)
def start(tag,title):
 n=len(p.slides)+1;s=p.slides.add_slide(p.slide_layouts[6]);rect(s,0,0,960,540,'bg');rect(s,42,30,26,4,'accent')
 txt(s,80,23,830,22,f'{n:02d} / {tag}',11,'accent',True);txt(s,42,65,880,88,title,34,bold=True)
 rect(s,42,505,876,1,'line');txt(s,42,518,800,14,'SMITHSONIAN × RICE D2K  ·  CLASSIFICATION UPDATE',9,'muted');txt(s,900,517,20,16,str(n),10,'muted')
 titles.append(title.replace('\n',' '));return s
def finish(s,note):
 notes.append(note);s.notes_slide.notes_text_frame.text=f'TARGET: {times[len(notes)-1]} seconds\n\n{note}';c.showPage()

s=start('PROJECT UPDATE','Two approaches to identifying fossil pollen')
txt(s,44,127,860,32,'RF-DETR method first · classifier backbones second',21,'muted')
for y,number,title,body in [(193,'01','RF-DETR','Find specimens and identify categories.'),(316,'02','Classifier backbones','Identify categories from specimen crops.')]:
 rect(s,44,y,374,100,'white');txt(s,59,y+14,40,29,number,18,'teal',True)
 txt(s,105,y+12,300,34,title,24,bold=True);txt(s,59,y+62,344,25,body,16,'muted')
for i,(name,label) in enumerate([('hero_0.png','Bisaccate'),('hero_1.png','Bombacacidites sp.'),('hero_2.png','Arecipites sp.')]):
 x=454+i*155;rect(s,x,214,145,190,'white');pic(s,name,x+6,221,133,143);txt(s,x+7,374,139,26,label,11,'muted')
txt(s,44,450,868,29,'Classification dataset: 2,283 specimens · 28 broad categories',18,'teal',True)
finish(s,"We’ve organized this update into two parts. First is the RF-DETR method, which aims to locate specimens and identify their categories together. Then we’ll cover the classifier backbones, which identify categories from prepared specimen crops. We kept the same 2,283 specimens and grouped the original 32 labels into 28 broader categories. Four Momipites labels now form one group, and two Caryapollenites labels form another. The images here are real examples with expert labels.")

s=start('SECTION 01 · RF-DETR','RF-DETR: detection and category identification')
rect(s,44,187,874,222,'white')
txt(s,79,267,800,48,'RF-DETR team update',32,'muted',True)
txt(s,79,329,800,30,'Placeholder · content to be added',20,'muted')
finish(s,"PLACEHOLDER — not a spoken script. RF-DETR team to add its method, progress and results. Allow approximately 60 seconds if the full presentation should stay near five minutes. No RF-DETR results are supplied or implied here.")

s=start('SECTION 02 · CLASSIFIER BACKBONES / FOCUS','A clearer view helps the model identify pollen')
for x,name,label in [(44,'focus_before.png','Original focal plane'),(302,'focus_after.png','Selected sharper plane')]:
 txt(s,x,158,248,30,label,18,bold=True);pic(s,name,x,198,230,230)
txt(s,44,451,510,30,'Same specimen and crop · expert label: Quercus sp.',14,'muted')
rect(s,577,166,341,276,'pale');txt(s,601,190,295,32,'Unseen-slide accuracy',19,'teal',True)
txt(s,601,244,292,68,'32% → 38%',43,'teal',True)
txt(s,601,331,295,82,'Earlier 32-category trial.\nSame Swin settings;\nselected plane changed.',19)
finish(s,"The microscope scans show each specimen at several focal depths. On the left is the original view; on the right is a sharper view of the same specimen. We used a sharpness score to choose one plane for each crop. In our earlier 32-category experiment, with the same Swin training settings, validation accuracy on unfamiliar slides increased from about 32 to 38 percent. This example illustrates the focus change; it doesn’t mean this particular specimen caused that improvement. We’d also like to know when identification requires features visible across several depths, rather than just one sharp image.")

s=start('CLASSIFIER BACKBONES / EVALUATION','Two questions, two ways to check performance')
for x,heading,sub in [(44,'Unseen slides','Can it work on a different slide?'),(492,'Shared slides','Can it identify more grains on familiar slides?')]:
 rect(s,x,164,424,283,'white');txt(s,x+20,183,390,34,heading,25,'teal',True);txt(s,x+20,226,388,25,sub,15,'muted')
 txt(s,x+20,280,160,25,'LEARN FROM',12,'muted',True);txt(s,x+230,280,170,25,'CHECK ON',12,'muted',True)
for x,labels in [(64,['Slide A','Slide B']),(274,['Slide C','Slide D']),(512,['A · grain 1','B · grain 1']),(722,['A · grain 2','B · grain 2'])]:
 for j,label in enumerate(labels):
  rect(s,x,315+j*48,168,36,'pale' if x in [274,722] else 'line');txt(s,x+12,323+j*48,147,25,label,16)
txt(s,44,467,874,25,'Different specimens in every set. Final test slides remain reserved.',18,'accent',True)
finish(s,"We checked performance in two situations. In the unseen-slide experiment, all specimens from a slide stay together, so the model is checked on slides it hasn’t learned from. In the shared-slide experiment, it learns from some grains and is checked on other grains from the same slides. We never reuse the exact same specimen in both sets. Shared slides can have familiar staining and imaging conditions, so that is an easier situation. Both experiments have the same number of training and validation examples, but the actual examples differ. A separate final test set remains unused.")

# Read the completed broad-category runs directly; preserve historical fine-label results.
def broad_metrics(model, split):
 sub='ordinary_baseline' if model=='swin' else 'backbone_finetune'
 path=BASE/f'broad_{model}_2026-10-08'/split/sub/'best_validation_metrics.json'
 return json.loads(path.read_text())
s=start('CLASSIFIER BACKBONES / RESULTS','DINO leads the broad-category comparison')
txt(s,42,133,880,26,'28 categories · validation accuracy on prepared specimen crops',19,'muted')
txt(s,315,186,245,30,'Unseen slides',23,bold=True);txt(s,646,186,260,30,'Shared slides',23,bold=True)
for i,(model,label) in enumerate([('swin','Swin-Tiny'),('dino','DINO · adapted')]):
 y=254+i*80;txt(s,43,y+2,249,35,label,23,bold=i==1)
 for x,split in [(315,'separate_slides'),(646,'shared_slides')]:
  v=100*broad_metrics(model,split)['accuracy']
  rect(s,x,y,200,28,'line');rect(s,x,y,200*v/100,28,'teal' if i==1 else 'gray');txt(s,x+210,y-2,80,36,f'{v:.1f}%',23,'teal' if i==1 else 'ink',True)
txt(s,44,416,874,25,'Grouping makes the task easier; retraining added little DINO accuracy.',18,'accent',True)
txt(s,44,450,870,25,'Earlier DINO predictions, regrouped: 50.1% unseen / 74.1% shared.',16,'muted')
txt(s,44,481,880,17,'Single-run validation comparisons; final test untouched. These are crop-classification results.',12,'muted')
finish(s,"We trained both models again using the 28 broader categories. Swin reached about 43 percent on unseen slides and 71 percent with shared slides. DINO reached 50 percent and 74 percent, so it still leads both comparisons. But grouping categories makes the task easier. If we simply regroup our earlier DINO predictions, we already get about 50 and 74 percent. So retraining on the broader labels has added very little accuracy so far. These are validation results for prepared crops; we have not evaluated the final test set.")

s=start('CLASSIFIER BACKBONES / ERRORS','Three categories still need closer review')
txt(s,44,132,874,30,'Adapted DINO · correctly identified / validation examples',19,'muted')
txt(s,44,190,425,24,'CATEGORY',12,'muted',True)
txt(s,520,185,190,30,'Shared slides',21,bold=True)
txt(s,733,185,190,30,'Unseen slides',21,bold=True)
lookup={split:{v['category']:v for v in broad_metrics('dino',split)['per_class']} for split in ['shared_slides','separate_slides']}
for i,name in enumerate(['Bombacacidites sp.','Polyatriopollenites type','Psilatricolpites sp.']):
 y=237+i*52;rect(s,44,y-6,874,45,'white');txt(s,55,y,440,32,name,22,bold=True)
 for x,split in [(520,'shared_slides'),(733,'separate_slides')]:
  v=lookup[split][name];correct=round(v['recall']*v['support'])
  txt(s,x,y,185,32,f"{correct}/{v['support']}  ({v['recall']:.0%})",22,'accent',True)
txt(s,44,403,874,26,'Fewer examples may contribute; the main cause is not yet established.',18,'accent',True)
txt(s,44,438,874,26,'Next: inspect the mistakes, image clarity, and confused categories.',18,'muted')
txt(s,44,478,874,20,'Small validation samples make these percentages uncertain. Momipites is no longer a main weakness.',12,'muted')
finish(s,"These three categories remain difficult in both settings: Bombacacidites, Polyatriopollenites, and Psilatricolpites. The counts show how many examples were identified correctly. They have relatively few labeled examples, which may contribute, but some other small categories perform well. We haven’t established the main cause. We’ll inspect which categories they get confused with and whether focus or crop quality is hiding useful details. The shared-slide results use only seven or eight examples per category, so the percentages are uncertain. The merged Momipites group now gets about 80 percent correct in both settings and is no longer a main weakness.")

p.save(OUT/'Smithsonian_sponsor_update.pptx');c.save()
(OUT/'speaker_notes.md').write_text('# Sponsor update speaker notes\n\nPrepared content: approximately 3:40. RF-DETR slide is an unfilled placeholder; allow about 60 seconds for that team’s update to target 4:40 total. No meeting date assumed.\n\n'+'\n\n'.join(f'## {i+1}. {t} — {times[i]} seconds\n\n{n}' for i,(t,n) in enumerate(zip(titles,notes))))
(OUT/'timing.json').write_text(json.dumps({'seconds':sum(times),'spoken_words':sum(len(n.split()) for n in notes),'slides':[{'slide':i+1,'title':t,'seconds':times[i],'words':len(notes[i].split())} for i,t in enumerate(titles)]},indent=2)+'\n')
print('Generated',len(p.slides),'slides;',sum(len(n.split()) for n in notes),'spoken words')
