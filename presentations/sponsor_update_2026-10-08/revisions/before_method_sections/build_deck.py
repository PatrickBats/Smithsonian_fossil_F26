from pathlib import Path
import json, textwrap, hashlib
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image

OUT=Path(__file__).resolve().parent;BASE=OUT.parents[2]
C={'bg':'F8F5EF','ink':'202D48','accent':'93451F','muted':'596477','line':'E6DFD4','white':'FFFFFF','teal':'20786E','pale':'E4EFEB','gray':'9AA5B4'}
pdfmetrics.registerFont(TTFont('Lato','/usr/share/fonts/truetype/lato/Lato-Regular.ttf'))
pdfmetrics.registerFont(TTFont('LatoB','/usr/share/fonts/truetype/lato/Lato-Bold.ttf'))
p=Presentation();p.slide_width=Inches(960/72);p.slide_height=Inches(540/72)
c=canvas.Canvas(str(OUT/'Smithsonian_sponsor_update.pdf'),pagesize=(960,540));c.setTitle('Smithsonian: classification progress')
notes=[];times=[40,45,50,55,50,40];titles=[]
def col(k):return '#'+C.get(k,k)
def rect(s,x,y,w,h,k):
 z=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));z.fill.solid();z.fill.fore_color.rgb=RGBColor.from_string(C.get(k,k));z.line.fill.background()
 c.setFillColor(col(k));c.rect(x,540-y-h,w,h,fill=1,stroke=0)
def txt(s,x,y,w,h,t,size=18,k='ink',bold=False):
 z=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));f=z.text_frame;f.word_wrap=False
 f.margin_top=f.margin_bottom=f.margin_left=f.margin_right=0
 for i,line in enumerate(t.split('\n')):
  q=f.paragraphs[0] if i==0 else f.add_paragraph();q.text=line;q.font.name='Lato';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(C.get(k,k));q.space_before=q.space_after=Pt(0);q.line_spacing=1.12
  assert pdfmetrics.stringWidth(line,'LatoB' if bold else 'Lato',size)<=w+4,(line,w)
  c.setFont('LatoB' if bold else 'Lato',size);c.setFillColor(col(k));c.drawString(x,540-y-size*.9-i*size*1.12,line)
 return z
def pic(s,name,x,y,w,h):
 path=OUT/'assets'/name
 with Image.open(path) as im:iw,ih=im.size
 ratio=min(w/iw,h/ih);ww,hh=iw*ratio,ih*ratio;x+=(w-ww)/2;y+=(h-hh)/2
 s.shapes.add_picture(str(path),Pt(x),Pt(y),Pt(ww),Pt(hh));c.drawImage(str(path),x,540-y-hh,width=ww,height=hh)
def start(tag,title):
 n=len(p.slides)+1;s=p.slides.add_slide(p.slide_layouts[6]);rect(s,0,0,960,540,'bg');rect(s,42,30,26,4,'accent')
 txt(s,80,23,830,22,f'{n:02d} / {tag}',11,'accent',True);txt(s,42,65,880,88,title,34,bold=True)
 rect(s,42,505,876,1,'line');txt(s,42,518,800,14,'SMITHSONIAN × RICE D2K  ·  CLASSIFICATION UPDATE',9,'muted');txt(s,900,517,20,16,str(n),10,'muted')
 titles.append(title.replace('\n',' '));return s
def finish(s,note):
 notes.append(note);s.notes_slide.notes_text_frame.text=f'TARGET: {times[len(notes)-1]} seconds\n\n{note}';c.showPage()

s=start('PROGRESS','From pollen images to category predictions')
txt(s,44,127,860,32,'A working baseline—and clearer questions for expert review.',21,'muted')
for x,num,label in [(44,'2,283','labeled specimens'),(240,'32','priority categories')]:
 txt(s,x,211,190,62,num,48,'teal',True);txt(s,x,276,190,30,label,17,'muted')
txt(s,44,353,340,85,'Prepared images\nCompared classifiers\nSaved results for review',21)
for i,(name,label) in enumerate([('hero_0.png','Bisaccate'),('hero_1.png','Bombacacidites sp.'),('hero_2.png','Arecipites sp.')]):
 x=454+i*155;rect(s,x,214,145,190,'white');pic(s,name,x+6,221,133,143);txt(s,x+7,374,139,26,label,11,'muted')
txt(s,454,431,456,28,'Real specimen crops · expert category labels',13,'muted')
finish(s,"Since our last update, we’ve moved from preparing the data to training initial classifiers. We now have 2,283 labeled specimens across 32 priority categories in this experiment. These are real examples from the collection, with the expert category labels. We prepared focused images, compared two model families, and saved their predictions for review. Our goal is still to help researchers identify specimens more efficiently. Today, we’ll show what improved, where the models still struggle, and where your input would be most useful.")

s=start('IMAGE PREPARATION','A clearer view helps the model identify pollen')
for x,name,label in [(44,'focus_before.png','Original focal plane'),(302,'focus_after.png','Selected sharper plane')]:
 txt(s,x,158,248,30,label,18,bold=True);pic(s,name,x,198,230,230)
txt(s,44,451,510,30,'Same specimen and crop · expert label: Quercus sp.',14,'muted')
rect(s,577,166,341,276,'pale');txt(s,601,190,295,32,'Unseen-slide accuracy',19,'teal',True)
txt(s,601,244,292,68,'32% → 38%',43,'teal',True)
txt(s,601,331,295,82,'Same Swin model settings.\nOnly the selected focal\nplane changed.',19)
finish(s,"The microscope scans show each specimen at several focal depths. On the left is the original view; on the right is a sharper view of the same specimen. We used a sharpness score to choose one plane for each crop. With the same Swin training settings, validation accuracy on unfamiliar slides increased from about 32 to 38 percent. This example illustrates the focus change; it doesn’t mean this particular specimen caused that improvement. We’d also like to know when identification requires features visible across several depths, rather than just one sharp image.")

s=start('EVALUATION','Two questions, two ways to check performance')
for x,heading,sub in [(44,'Unseen slides','Can it work on a different slide?'),(492,'Shared slides','Can it identify more grains on familiar slides?')]:
 rect(s,x,164,424,283,'white');txt(s,x+20,183,390,34,heading,25,'teal',True);txt(s,x+20,226,388,25,sub,15,'muted')
 txt(s,x+20,280,160,25,'LEARN FROM',12,'muted',True);txt(s,x+230,280,170,25,'CHECK ON',12,'muted',True)
for x,labels in [(64,['Slide A','Slide B']),(274,['Slide C','Slide D']),(512,['A · grain 1','B · grain 1']),(722,['A · grain 2','B · grain 2'])]:
 for j,label in enumerate(labels):
  rect(s,x,315+j*48,168,36,'pale' if x in [274,722] else 'line');txt(s,x+12,323+j*48,147,25,label,16)
txt(s,44,467,874,25,'Different specimens in every set. Final test slides remain reserved.',18,'accent',True)
finish(s,"We checked performance in two situations. In the unseen-slide experiment, all specimens from a slide stay together, so the model is checked on slides it hasn’t learned from. In the shared-slide experiment, it learns from some grains and is checked on other grains from the same slides. We never reuse the exact same specimen in both sets. Shared slides can have familiar staining and imaging conditions, so that is an easier situation. Both experiments have the same number of training and validation examples, but the actual examples differ. A separate final test set remains unused.")

s=start('RESULTS','Adapting DINO gives our strongest results so far')
txt(s,42,133,880,26,'Validation accuracy · identifying categories in prepared specimen crops',17,'muted')
txt(s,315,186,245,30,'Unseen slides',23,bold=True);txt(s,646,186,260,30,'Shared slides',23,bold=True)
models=[('Swin',38.0403458213,66.2824207493),('DINO · frozen',29.9711815562,59.0778097983),('DINO · adapted',43.5158501441,70.0288184438)]
for i,(label,left,right) in enumerate(models):
 y=249+i*66;txt(s,43,y+2,249,35,label,23,bold=i==2)
 for x,v in [(315,left),(646,right)]:
  rect(s,x,y,210,28,'line');rect(s,x,y,210*v/100,28,'teal' if i==2 else 'gray');txt(s,x+220,y-2,66,36,f'{v:.0f}%',25,'teal' if i==2 else 'ink',True)
txt(s,44,454,870,25,'Adapted = trained the final two DINO blocks and its category classifier.',16,'muted')
txt(s,44,481,880,17,'Single-run validation comparisons; final test untouched. These are not whole-slide detection results.',11,'muted')
finish(s,"Here are the classification results. Swin reached 38 percent on unseen slides and 66 percent with shared slides. DINO initially did worse when we kept its image-processing backbone fixed and trained only a small classifier. We then allowed the final part of DINO to adapt to our specimens. That reached about 44 percent on unseen slides and 70 percent with shared slides. Its category-averaged F1 score also improved. These are promising initial results, but they are single-run validation comparisons, not final test results. They measure identification of prepared crops, rather than the complete process of finding and classifying everything on a slide.")

s=start('ERROR REVIEW','Which differences matter to an expert?')
assets=json.loads((OUT/'asset_manifest.json').read_text())
for i in range(3):
 a=next(a for a in assets if a['file']==f'error_{i}.png');x=44+i*296;rect(s,x,155,278,311,'white');pic(s,a['file'],x+44,166,190,190)
 txt(s,x+15,368,250,18,'EXPERT LABEL',10,'muted',True)
 label=a['category'].replace('Momipites wyomingensis','Momipites\nwyomingensis')
 txt(s,x+15,388,249,44,label,17,bold=True)
 txt(s,x+15,441,250,23,'Model: '+a['predicted_category'],14,'accent')
txt(s,44,477,878,22,'Selected validation mistakes—not proposed changes to the expert labels.',14,'muted')
finish(s,"These are actual mistakes from the adapted DINO model on unseen-slide validation specimens. The expert labels are shown above the model’s predictions. We selected examples from recurring confusions, rather than treating them as representative of every error. We don’t yet know whether these mistakes reflect subtle biological differences, image quality, nearby material, or the model relying on the wrong features. Your interpretation would help us decide what to improve. We would especially value guidance on which visible characteristics distinguish these categories, and whether another focal view would make the distinction clearer.")

s=start('DISCUSSION','Help us prioritize the next improvements')
questions=['Which category confusions matter most scientifically?', 'Which categories should we improve first?', 'Can we review a small set of confusing specimens together?', 'When is more than one focal view essential?']
for i,q in enumerate(questions):
 y=163+i*64;txt(s,46,y,40,32,f'0{i+1}',19,'teal',True);txt(s,100,y,809,40,q,22)
rect(s,44,438,874,49,'pale');txt(s,60,450,842,29,'Next: expert error review → targeted improvements → final evaluation',19,'teal',True)
finish(s,"The main takeaway is that we now have a working classifier and a way to compare improvements. Adapting DINO helped, but unfamiliar slides remain challenging. We’d like your help deciding which confusions have the greatest scientific consequences, which categories deserve priority, and when multiple focal views are essential. Our next step is to review a manageable set of errors with you, make targeted improvements, and then evaluate the selected approach on the reserved test slides. We want the next experiments to address useful scientific distinctions, not simply produce a higher overall score.")

p.save(OUT/'Smithsonian_sponsor_update.pptx');c.save()
(OUT/'speaker_notes.md').write_text('# Sponsor update speaker notes\n\nTarget: approximately 4:40 speaking, plus transitions. No meeting date assumed.\n\n'+'\n\n'.join(f'## {i+1}. {t} — {times[i]} seconds\n\n{n}' for i,(t,n) in enumerate(zip(titles,notes))))
(OUT/'timing.json').write_text(json.dumps({'seconds':sum(times),'spoken_words':sum(len(n.split()) for n in notes),'slides':[{'slide':i+1,'title':t,'seconds':times[i],'words':len(notes[i].split())} for i,t in enumerate(titles)]},indent=2)+'\n')
print('Generated',len(p.slides),'slides;',sum(len(n.split()) for n in notes),'spoken words')
