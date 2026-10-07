"""Build an editable replacement for the historical category inventory slide."""
from pathlib import Path
import csv
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import fitz
P=Path(__file__).resolve().parent
ROOT=P.parents[1]
rows=list(csv.DictReader((ROOT/'data/current/category_summary.csv').open()))
a=sorted([r for r in rows if r['priority']=='YES'],key=lambda r:-int(r['annotation_records']))
assert len(a)==32 and sum(int(r['annotation_records']) for r in a)==2290
assert all(int(r['difference'])==0 for r in rows)
missing={'Momipites flexus','Momipites ventifluminis','Nudopollis sp.','Psilatricolpites sp.','Rousea sp.'}
for name,file in [('Lato','Lato-Regular.ttf'),('Lato-Bold','Lato-Bold.ttf')]:
 pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/lato/'+file))
prs=Presentation();prs.slide_width=Pt(960);prs.slide_height=Pt(540)
s=prs.slides.add_slide(prs.slide_layouts[6])
pdf=canvas.Canvas(str(P/'Updated_category_inventory.pdf'),pagesize=(960,540))
colors={'bg':'F6F4EE','ink':'142F38','muted':'596D71','teal':'16756D','light':'E3EEEA','white':'FFFFFF'}
def rect(x,y,w,h,c):
 c=colors.get(c,c);sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background()
 pdf.setFillColor('#'+c);pdf.rect(x,540-y-h,w,h,stroke=0,fill=1)
def text(t,x,y,w,size=14,color='ink',bold=False):
 c=colors.get(color,color);font='Lato-Bold' if bold else 'Lato'
 assert pdfmetrics.stringWidth(t,font,size)<=w,(t,w)
 sh=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(size*1.5));tf=sh.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 p=tf.paragraphs[0];p.text=t;p.font.name='Lato';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(c)
 pdf.setFont(font,size);pdf.setFillColor('#'+c);pdf.drawString(x,540-y-size,t)
rect(0,0,960,540,'bg')
text('01/Data update',36,23,700,11,'teal',True)
text('Updated category mapping',36,47,895,28,bold=True)
text('2,290 target annotations  •  83 image/annotation pairs in the full dataset',36,89,895,17,'muted')
for col in range(2):
 x=36+col*454
 text('CATEGORY',x+8,129,300,11,'muted',True)
 text('COUNT',x+330,129,60,11,'muted',True)
 text('SLIDES',x+393,129,52,11,'muted',True)
 for j,r in enumerate(a[col*16:(col+1)*16]):
  y=151+j*19
  rect(x,y-1,440,19,'white')
  text(r['category'],x+8,y,317,13,bold=False)
  text(r['annotation_records'],x+339,y,48,13,'teal',True)
  text(r['slides'],x+402,y,35,13)
text('All 185 title counts and 65 category totals match the sponsor tables.',36,463,895,14,'teal',True)
text('Full dataset: 11,275 named annotations • 65 categories • 185 annotation titles',36,492,895,13,'muted')
notes='''This replaces the old category inventory slide, which listed 27 represented YES categories and 933 annotations. The corrected Fall 2026 annotation set now has 2,290 records across all 32 YES categories. Five categories were missing in the earlier inventory: Nudopollis sp. (89), Momipites ventifluminis (61), Psilatricolpites sp. (61), Rousea sp. (54), and Momipites flexus (38).

COUNT means annotation records; SLIDES is the number of different slide files containing that category. Categories are sponsor groupings, not necessarily individual species. These are inventory counts before crop-quality review, sampling or train/test splitting. They do not establish that every record is a unique usable training specimen.

The complete set has 83 verified image/annotation pairs, 11,275 named records, 185 annotation titles and 65 categories. The breakdown is 32 YES / 2,290 records; 5 MAYBE / 114 records; 28 NO / 8,871 records. Another 227 untitled marks are outside these named counts. Every supplied title and category count reconciles after trimming trailing spaces in three labels. Provisional biological aliases from the old inventory are no longer needed.

Source: data/current/category_summary.csv, title_reconciliation.csv, audit_manifest.json and cluster_verification.json. Corrected source package: Fall_2026.zip received September 21. Slide prepared September 25, 2026.
'''
s.notes_slide.notes_text_frame.text=notes
prs.save(P/'Updated_category_inventory.pptx');pdf.save()
(P/'speaker_notes.txt').write_text(notes)
doc=fitz.open(P/'Updated_category_inventory.pdf');doc[0].get_pixmap(matrix=fitz.Matrix(2,2)).save(P/'Updated_category_inventory.png')
print('Created PPTX, PDF, PNG and speaker notes.')
