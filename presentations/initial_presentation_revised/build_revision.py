from pathlib import Path
from io import BytesIO
import hashlib,json,re
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
OUT=Path(__file__).resolve().parent
SRC=Path('/home/pb52/.codex/attachments/6eb22bcf-a1da-4d6a-91f4-d22aa2fd0363/Initial Presentation.pptx')
p=Presentation(SRC)
C={'bg':'F8F5EF','ink':'202D48','accent':'93451F','light':'EEE3D6','muted':'596477','white':'FFFFFF'}
def text(s,x,y,w,h,t,size=16,color='ink',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line;q.font.name='Lato';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(C[color]);q.space_after=Pt(5)
 return sh
def rect(s,x,y,w,h,c):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C[c]);sh.line.fill.background();return sh
def walk(ss):
 for sh in ss:
  if sh.shape_type==6: yield from walk(sh.shapes)
  else: yield sh
def settext(sh,t,size=None):
 from copy import deepcopy
 tf=sh.text_frame
 q=tf.paragraphs[0]
 rp=deepcopy(q.runs[0]._r.rPr) if q.runs and q.runs[0]._r.rPr is not None else None
 pp=deepcopy(q._p.pPr) if q._p.pPr is not None else None
 tf.clear()
 for i,line in enumerate(t.split('\n')):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph()
  if pp is not None:
   if q._p.pPr is not None:q._p.remove(q._p.pPr)
   q._p.insert(0,deepcopy(pp))
  r=q.add_run();r.text=line
  if rp is not None:r._r.insert(0,deepcopy(rp))
  if size:r.font.size=Pt(size)
 tf.word_wrap=True
 return sh
def byid(s,n):return next(sh for sh in walk(s.shapes) if sh.shape_id==n)
def change(s,n,t,size=None):return settext(byid(s,n),t,size)
def clean(s):
 for sh in list(s.shapes):sh._element.getparent().remove(sh._element)
def frame(s,num,section,title):
 rect(s,0,0,10,5.625,'bg');rect(s,.38,.29,.26,.04,'accent')
 text(s,.76,.20,8.85,.2,f'{num:02d} / {section}',8,'accent',True)
 if title:text(s,.38,.59,9.25,.6,title,24,bold=True)
 rect(s,.38,5.26,9.25,.009,'light');text(s,.38,5.36,8.65,.18,'SMITHSONIAN × RICE D2K',7,'muted');text(s,9.35,5.34,.3,.2,str(num),8,'muted')
# Existing background: keep its three-card layout, simplify claims and acronyms.
s=p.slides[0]
change(s,92,'Why fossil pollen matters',24)
change(s,97,'Pollen survives in sediment.\nIt records ancient plant life.\nWe study vegetation about 50 million years ago.',14)
change(s,101,'The challenge',17)
change(s,102,'Microscope images contain many specimens and debris.\nIdentifying each specimen takes expert time.',14)
change(s,106,'Our starting point',17)
change(s,107,'The previous Rice team developed a detector.\nOur next step is identifying the category of each specimen.',14)
# Three related-work cards; the third paper supplies whole-slide detection.
s=p.slides[2];clean(s);frame(s,4,'PRIOR WORK','What earlier studies suggest')
for x,num,author,body in [(.4,'01','Punyasena et al.','Locate pollen\n↓\nCrop specimens\n↓\nIdentify categories'),(3.55,'02','Martinsen et al.','Locate specimens\n↓\nCrop specimens\n↓\nIdentify categories'),(6.7,'03','Abbas Shaikh et al.','Multifocal slides\n↓\nImage tiles\n↓\nDetect specimens')]:
 rect(s,x,1.42,2.9,2.98,'white')
 text(s,x+.18,1.64,2.54,.25,num,12,'accent',True)
 text(s,x+.18,2.03,2.54,.37,author,17,bold=True)
 box=text(s,x+.18,2.64,2.54,1.65,body,15,'muted')
 for para in box.text_frame.paragraphs:para.alignment=PP_ALIGN.CENTER;para.space_after=Pt(1)
text(s,.4,4.75,9.2,.32,'Build on whole-slide detection; add category identification.',16,'accent',True)
# Example-specific microscopy details.
s=p.slides[3];change(s,154,'One specimen, multiple focal views',24)
old=byid(s,157);old._element.getparent().remove(old._element)
text(s,6.37,1.29,2.85,.65,'This example: 25 focal planes\nAbout 30 GB for the full scan',12,'muted')
# Dataset metrics: improve scannability without changing the original chart.
s=p.slides[4]
old=byid(s,170);old._element.getparent().remove(old._element)
for y,value,label in [(1.43,'83','image–annotation pairs'),(2.39,'32','priority categories'),(3.35,'2,290','annotations')]:
 text(s,.4,y,2.5,.48,value,28,'ink',True)
 text(s,.4,y+.5,2.65,.27,label,12,'muted')
text(s,.4,4.52,2.7,.47,'Some categories have\nfar fewer examples.',12,'accent',True)
# Simplify data preparation language.
s=p.slides[5];change(s,180,'Connect images, locations, and expert labels',24)
change(s,186,'Match the files',18);change(s,194,'Check the examples',18)
change(s,195,'Check positions, missing\nlabels, and duplicates.',14)
change(s,199,'Each specimen linked to its image, location, and category',14)
s=p.slides[6];change(s,208,'Prepare examples for a fair test',24)
change(s,219,'Extract small image regions;\ncheck position and focus.',14)
change(s,223,'Regions for detection;\nspecimen crops for\nclassification.',14)
change(s,227,'Vary training examples and balance categories.',14)
# Rebuild only the crowded methods slide; preserve its source images in the file.
s=p.slides[7]
method_images=[sh.image.blob for sh in s.shapes if sh.shape_type==13]
clean(s);frame(s,9,'PROPOSED METHODS','Two ways to identify specimens')
text(s,.4,1.32,9,.3,'1   FIND AND IDENTIFY TOGETHER',13,'accent',True)
rect(s,.4,1.85,4.1,.8,'white');text(s,.6,2.02,3.7,.45,'RF-DETR',20,bold=True)
text(s,4.65,2.02,.5,.4,'→',23,'accent')
rect(s,5.25,1.85,4.35,.8,'light');text(s,5.45,2.04,3.95,.45,'Location + category',19,bold=True)
text(s,.4,3.0,9,.3,'2   FIND FIRST, IDENTIFY SECOND',13,'accent',True)
for x,w,head,body in [(.4,2.65,'RF-DETR','Find specimens'),(3.65,2.65,'Specimen crops','Isolate each example'),(6.9,2.7,'Swin-Tiny','Identify its category')]:
 rect(s,x,3.52,w,.97,'white');text(s,x+.16,3.65,w-.32,.3,head,18,bold=True);text(s,x+.16,4.05,w-.32,.28,body,12,'muted')
for x in [3.15,6.4]:text(s,x,3.82,.4,.4,'→',23,'accent')
text(s,.4,4.78,9.2,.3,'Compare location accuracy and category identification.',15,'accent',True)
# Welcome and conclusion as editable slides.
w=p.slides.add_slide(p.slide_layouts[6]);frame(w,1,'WELCOME','')
rect(w,.4,1.35,.055,2.3,'accent')
text(w,.68,1.37,6.2,1.65,'Identifying Fossil Pollen\nwith Machine Learning',30,bold=True)
text(w,.7,3.15,5.7,.65,'Help researchers identify fossil pollen\ncategories more efficiently.',17,'muted')
fig=OUT.parents[1]/'reports/initial_report_complete_2026-09-27/figures/specimen_sources/shard_0_target_027_crop.png'
w.shapes.add_picture(str(fig),Inches(7.05),Inches(1.53),width=Inches(2.45),height=Inches(2.45))
text(w,.7,4.48,8.7,.4,'Patrick Batsell · Yun Ying Tsai · Yeonju Kim · Yunfan Bao · Alan Yang',12,'muted')
ids=p.slides._sldIdLst;el=ids[-1];ids.remove(el);ids.insert(0,el)
s=p.slides.add_slide(p.slide_layouts[6]);frame(s,10,'CONCLUSION','From finding pollen to identifying it')
text(s,.45,1.52,8.9,.6,'Compare two approaches to automated identification.',23,bold=True)
for x,n,title in [(.45,'01','Prepare specimen images'),(3.6,'02','Train initial models'),(6.75,'03','Evaluate mistakes')]:
 rect(s,x,2.55,2.8,1.28,'white');text(s,x+.18,2.73,2.45,.25,n,12,'accent',True);text(s,x+.18,3.12,2.43,.55,title,17,bold=True)
text(s,.45,4.35,9,.55,'Help experts study past vegetation more efficiently.',21,'accent',True)
# Standardize all retained slide headers and footers.
sections=['WELCOME','BACKGROUND','THE PROBLEM','PRIOR WORK','MICROSCOPY DATA','ANNOTATION INVENTORY','DATA PIPELINE','DATA PIPELINE','PROPOSED METHODS','CONCLUSION']
for i,s in enumerate(p.slides):
 if i in [0,8,9]:continue
 for sh in s.shapes:
  if not sh.has_text_frame:continue
  if sh.top<Inches(.45) and ' / ' in sh.text:settext(sh,f'{i+1:02d} / {sections[i]}')
  elif sh.top>Inches(5.3):
   if 'SMITHSONIAN' in sh.text:settext(sh,'SMITHSONIAN × RICE D2K')
   elif sh.text.strip().isdigit():settext(sh,str(i+1))
notes=[
"Hi everyone. We’re the Rice D2K Smithsonian team. Our project is about helping researchers identify fossil pollen using machine learning.",
"Fossil pollen gives us clues about which plants lived in the past. Smithsonian researchers use it to understand North American vegetation about fifty million years ago. The challenge is that identifying these tiny specimens takes expert time. The previous Rice team developed a detector that finds specimens. We want to build on that by identifying their categories.",
"This shows the difference between finding a specimen and identifying it. On the left, the box marks one pollen grain in a crowded microscope image. On the right, we can inspect its shape and surface details. Detection asks, where is it? Classification asks, what kind is it? Our project connects those two tasks.",
"These studies give us two foundations. Punyasena and Martinsen locate specimens and then classify individual crops. Abbas Shaikh and colleagues developed a pipeline for detecting specimens across large, multifocal slides. We want to build on that detection work and add category identification, comparing separate and combined approaches.",
"The images also capture different focal depths. Here, the same specimen appears across twenty-five focal planes, and this example’s full scan is about thirty gigabytes. We’ll work with smaller image regions and use the focal views to make the specimen’s details easier to see.",
"Our collection has eighty-three microscope images paired with expert annotations. We’re starting with thirty-two priority categories containing two thousand two hundred ninety annotations. The chart shows that some categories have many more examples than others. That matters because categories with fewer examples can be harder to learn. We’ll account for that when training and evaluating the models.",
"First, we organize the examples. Each microscope image is matched with its annotation file, which gives us specimen locations and expert labels. We connect those labels to the categories we want to recognize. Then we check that the marks point to the right specimens and look for missing labels or duplicates. This keeps each example linked to its source.",
"Next, we separate the data for training, validation, and testing. Examples from the same microscope slide stay together, so the test checks performance on unfamiliar slides. We then extract smaller regions and check their position and focus. Larger regions help the detector find specimens, while individual crops help the classifier identify them. Any extra image variations or category balancing are applied only to training examples.",
"We’ll compare two approaches. The first uses RF-DETR to find each specimen and predict its category at the same time. The second uses RF-DETR just to find specimens. We then crop each specimen and pass it to Swin-Tiny, a separate classifier. This comparison helps us see whether identifying specimens separately works better than doing both tasks together. We’ll check whether the models find the right locations and assign the right categories, including categories with fewer examples. These are our proposed approaches; training and evaluation are the next steps.",
"The main idea is to move from finding pollen to identifying it. Next, we’ll prepare the specimen images, train the initial models, and look closely at their mistakes. The goal is to give experts a useful starting point for identification, so they can study past vegetation more efficiently. Thank you."
]
times=[10,30,25,25,20,30,30,35,50,25]
titles=['Welcome','Background','The problem','Previous approaches','Microscopy data','Annotation inventory','Organizing the data','Preparing examples','Proposed methods','Conclusion']
assert len(p.slides)==10
manifest=[];md=['# Speaker notes\n','Target: 4:40 speaking, with up to 20 seconds for transitions. The rubric specifies exactly five minutes; rehearse to approach five minutes without running over.\n','Presenter blocks: 1 → slides 1–2; 2 → slides 3–4; 3 → slides 5–6; 4 → slides 7–8; 5 → slides 9–10. Assign names within the team.\n']
for i,(s,n,t,title) in enumerate(zip(p.slides,notes,times,titles)):
 count=len(n.split());s.notes_slide.notes_text_frame.text=f'TARGET: {t} seconds | PRESENTER {i//2+1}\n\nSAY:\n{n}'
 manifest.append({'slide':i+1,'title':title,'seconds':t,'words':count,'presenter':i//2+1})
 md.append(f'## Slide {i+1}: {title} — {t} seconds\n\n{n}\n')
assert 530<=sum(x['words'] for x in manifest)<=560,sum(x['words'] for x in manifest)
p.save(OUT/'Smithsonian_initial_presentation_revised.pptx')
(OUT/'speaker_notes.md').write_text('\n'.join(md))
(OUT/'timing.json').write_text(json.dumps(manifest,indent=2)+'\n')
(OUT/'source_manifest.json').write_text(json.dumps({'source_filename':SRC.name,'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'slides':10,'spoken_words':sum(x['words'] for x in manifest),'planned_seconds':sum(times),'original_preserved':True},indent=2)+'\n')
print('WORDS',sum(x['words'] for x in manifest));print(manifest)
