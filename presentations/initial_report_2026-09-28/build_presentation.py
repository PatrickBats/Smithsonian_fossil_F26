from pathlib import Path
import textwrap,json,hashlib,re
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image
import fitz
P=Path(__file__).resolve().parent;R=P.parents[1];A=R/'reports/initial_report_complete_2026-09-27/figures/specimen_sources'
W,H=960,540
C={'bg':'F8F5EF','ink':'202D48','teal':'93451F','light':'EEE3D6','muted':'596477','white':'FFFFFF','gold':'B97832'}
for n,f in [('Lato','Lato-Regular.ttf'),('Lato-Bold','Lato-Bold.ttf')]:pdfmetrics.registerFont(TTFont(n,'/usr/share/fonts/truetype/lato/'+f))
prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
pdf=canvas.Canvas(str(P/'Smithsonian_initial_presentation.pdf'),pagesize=(W,H));notes=[];bounds=[]
def rect(x,y,w,h,c):
 c=C.get(c,c);sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();pdf.setFillColor('#'+c);pdf.rect(x,H-y-h,w,h,stroke=0,fill=1)
def tx(t,x,y,w,size=20,c='ink',bold=False):
 font='Lato-Bold' if bold else 'Lato';lines=[]
 for para in t.split('\n'):
  line=''
  for word in para.split():
   cand=(line+' '+word).strip()
   if line and pdfmetrics.stringWidth(cand,font,size)>w:lines.append(line);line=word
   else:line=cand
  lines.append(line)
 height=len(lines)*size*1.2+4;bounds.append((len(prs.slides),t[:40],x,y,w,height))
 sh=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(height));tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Lato';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(c,c));p.line_spacing=Pt(size*1.2);p.space_after=Pt(0)
 pdf.setFont(font,size);pdf.setFillColor('#'+C.get(c,c))
 for i,line in enumerate(lines):pdf.drawString(x,H-y-size-i*size*1.2,line)
def pic(path,x,y,w,h):
 iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale;x+=(w-ww)/2;y+=(h-hh)/2
 s.shapes.add_picture(str(path),Pt(x),Pt(y),width=Pt(ww),height=Pt(hh));pdf.drawImage(str(path),x,H-y-hh,ww,hh)
def new(k,title):
 global s
 s=prs.slides.add_slide(prs.slide_layouts[6]);rect(0,0,W,H,'bg');rect(36,28,25,4,'teal');tx(k,73,20,850,11,'teal',True);tx(title,36,57,888,32,bold=True)
 rect(36,505,888,1,'light');tx('SMITHSONIAN × RICE D2K  /  INITIAL PROJECT PRESENTATION',36,516,830,9,'muted');tx(str(len(prs.slides)),900,514,24,11,'teal',True)
def end(title,seconds,spoken,source):
 n=f'TARGET: {seconds} seconds\n\nSAY:\n{spoken}\n\nSOURCE / PRESENTER CONTEXT (not spoken):\n{source}'
 s.notes_slide.notes_text_frame.text=n;notes.append(dict(slide=len(prs.slides),title=title,seconds=seconds,script=spoken,source=source));pdf.showPage()
def banner(t):rect(36,444,888,43,'light');tx(t,50,454,860,18,'teal',True)
new('01 / THE PROJECT','Fossil palynomorph detection and classification')
tx('One joint model,\nor a detector plus\na specialist classifier?',36,151,480,33,bold=True)
pic(A/'shard_0_target_027_crop.png',579,142,290,270)
tx('Patrick Batsell • Yun Ying Tsai • Yeonju Kim\nYunfan Bao • Alan Yang',36,363,510,16,'muted')
end('Project and objective',20,'Our project asks whether fossil specimens are better located and identified by one joint model, or by a detector followed by a separate classifier. We are building on existing Smithsonian and Rice work to support expert review of fossil pollen and other palynomorphs.', 'Report pp. 3–4, objectives. Palynomorphs include pollen and spores. Image: expert-labeled Bombacacidites focus-stacked crop.')
new('02 / SCIENTIFIC MOTIVATION','A vast collection, examined one specimen at a time')
for x,num,lab in [(36,'~50 million','years into the past'),(490,'~70,000','slides in the broader collection')]:
 rect(x,148,434,159,'white');tx(num,x+22,167,390,43,'teal',True);tx(lab,x+22,242,390,21,'muted')
tx('Pollen helps reconstruct vegetation and warm climates.',36,347,888,27,bold=True)
banner('The challenge is scaling expert identification.')
end('Motivation and scale',25,'Fossil pollen provides evidence about past vegetation and warm climates around fifty million years ago. The report describes a broader Smithsonian collection of roughly seventy thousand slides, far larger than our working dataset. Locating and identifying specimens manually is time-consuming. Automation could help scientists review more material while retaining access to the original images.', 'Report p. 3, Introduction; collection estimate attributed there to Romero et al., report reference [2]. This collection estimate is distinct from the 83-slide working dataset.')
new('03 / IMAGE CHALLENGES','Important features are small—and not always in focus')
for x,file,label in [(36,'shard_0_target_008_crop.png','Subtle morphology'),(338,'shard_0_target_027_crop.png','Staining differences'),(640,'shard_0_target_000_crop.png','Debris and overlap')]:
 pic(A/file,x,142,284,222);tx(label,x,381,284,23,bold=True)
banner('Multiple focal depths capture different parts of each specimen.')
end('Why recognition is difficult',25,'Identification depends on features such as shape, apertures, and surface structure. Different structures may be sharp at different focal depths, while staining and surrounding debris change the appearance. These real examples illustrate that variation. Focus stacking produces a usable two-dimensional image, but we must check whether it preserves the details needed for classification.', 'Report pp. 3–5, multifocal imaging and data quality; p. 7 specimen examples. These are different specimens with different source field widths, not a focal-plane sequence or a controlled stain comparison.')
new('04 / RELATED WORK','Build on detection; test how best to add identification')
for i,(head,body) in enumerate([('Whole-slide detection','Shaikh et al.: tiling, focal compression, and detection'),('Separate crop classification','Punyasena et al. and Martinsen et al.: two-stage precedents'),('Our comparison','Joint RF-DETR versus RF-DETR + Swin-Tiny')]):
 y=145+i*91;tx(str(i+1).zfill(2),36,y,50,23,'teal',True);tx(head,109,y,815,25,bold=True);tx(body,109,y+36,815,19,'muted')
end('Prior research and rationale',30,'Prior work gives us two foundations. Shaikh and colleagues developed whole-slide detection with image tiling and focal compression. Studies by Punyasena and Martinsen provide precedents for detecting specimens and classifying crops separately. Those studies use different images and categories, so their results do not establish performance here. Our contribution is to compare joint and two-stage identification on the same fossil dataset.', 'Report pp. 3–4, related work, references [4], [5], [7]. No claim that earlier studies used Swin or establish its superiority. Swin rationale: hierarchical local-window image features, report reference [9].')
new('05 / THE DATA','83 annotated slides; uneven category coverage')
for x,num,label in [(36,'11,275','named annotations overall'),(338,'2,290','initial target annotations'),(640,'32','candidate target categories')]:
 rect(x,136,284,121,'white');tx(num,x+18,149,246,41,'teal',True);tx(label,x+18,215,250,17,'muted')
for y,label,n in [(288,'Bisaccate',316),(352,'Cupuliferoipollenites sp.',27)]:
 tx(label,36,y,390,20,bold=True);rect(426,y+9,385,16,'light');rect(426,y+9,385*n/316,16,'teal');tx(str(n),843,y-2,75,25,'teal',True)
banner('Training options: undersampling, balanced batches, or weighted loss.')
end('Data and imbalance',30,'The working dataset contains eighty-three annotated slides and eleven thousand two hundred seventy-five named annotations. Our initial target inventory uses two thousand two hundred ninety annotations across thirty-two candidate categories. Counts range from twenty-seven to three hundred sixteen per target category. We will assess both category frequency and independent slide coverage, and consider training-only balancing rather than changing the evaluation data.', 'Report pp. 4–5, data inventory and imbalance. These are records before quality review, duplicate checking, regrouping and sampling. The two bars show the range on a common scale.')
new('06 / SHARED PREPROCESSING','Turn multifocal slides into traceable model inputs')
for x,n,h,b in [(36,'01','Define source groups','Separate slides and related samples before training.'),(338,'02','Prepare image tiles','Combine focal information; read overlapping regions.'),(640,'03','Preserve specimen identity','Retain coordinates, expert labels, and crop settings.')]:
 rect(x,153,284,244,'white');tx(n,x+18,169,246,18,'teal',True);tx(h,x+18,211,246,24,bold=True);tx(b,x+18,292,246,19,'muted')
banner('Same source groups, category mapping, and preprocessing for both models.')
end('Shared pipeline',25,'Both approaches share the same preparation and source groups. We will check annotation geometry, focus, and specimen quality, and keep related slides and all views of a specimen together when defining splits. Images are processed as overlapping, focus-stacked tiles. Every derived specimen retains its source coordinates and labels, so predictions can be traced back and inspected.', 'Report p. 5, Pipeline. Source groups assigned before tiling, augmentation or balancing; split proportions and exact settings remain to be specified.')
new('07 / APPROACH A','Joint localization and category identification')
pic(P/'assets/workflow_direct.png',36,118,888,318)
banner('Advantage: one model. Challenge: rare or similar categories.')
end('Multiclass RF-DETR',25,'The first approach adapts RF-DETR to predict specimen locations and categories together. It processes image tiles, and overlapping predictions are merged in whole-slide coordinates. This gives a unified detection-and-identification model. The question is whether its joint objective can learn uncommon or visually similar categories well enough.', 'Report pp. 6 and 8, Model 1. Architecture restored from report and recolored to match this deck. Model internals and output marks are conceptual, not measured predictions; thumbnails are real images.')
new('08 / APPROACH B','Separate localization from category identification')
pic(P/'assets/workflow_two_stage.png',36,118,888,318)
banner('Advantage: focused classification. Challenge: detection errors propagate.')
end('RF-DETR plus Swin-Tiny',25,'The second approach first uses RF-DETR to find specimens, then applies Swin-Tiny to their crops. This separates identification from localization and makes classifier errors easier to examine. However, missed detections or poor crops can reduce overall performance. Comparing expert-defined crops with detected crops will help identify that bottleneck.', 'Report pp. 7–8, Model 2. Same source split and preprocessing; detector and classifier scores retained separately. Proposed model, not completed classifier training.')
new('09 / EVALUATION','Test on separate source groups; report errors by category')
for x,h,b in [(36,'Detection','Missed specimens\nand false detections'),(338,'Classification','Macro-F1, balanced accuracy,\nand category-level errors'),(640,'Combined workflow','Specimens both found\nand correctly identified')]:
 rect(x,143,284,169,'white');tx(h,x+18,162,246,24,bold=True);tx(b,x+18,219,246,19,'muted')
rect(36,346,888,100,'light');tx('Initial YOLO check: 54/64 and 51/64 known targets recovered',51,358,858,22,'teal',True);tx('Feasibility only: centered crops and possible prior training exposure.',51,402,858,18,'muted')
end('Evaluation and preliminary evidence',35,'We will evaluate detection, classification, and the combined workflow on separate source groups, with per-category metrics so common categories do not dominate. Detection precision requires sufficiently complete annotations. As preliminary evidence, existing YOLO checkpoints recovered fifty-four and fifty-one of sixty-four known targets. These were centered crops with possible prior training exposure, so the results establish feasibility rather than independent generalization. The proposed RF-DETR and Swin comparison remains to be run.', 'Report pp. 8–9. Pilot: best-plane 54/64; focus-stack 51/64, confidence and IoU each >=0.50. Different checkpoints mean preprocessing is not isolated. Macro-F1 averages category-level F1; balanced accuracy averages recall.')
new('10 / NEXT STEPS','Compare the models, then make the results useful')
for x,n,t in [(36,'01','Inspect crops\nand define categories'),(338,'02','Freeze evaluation\nand train baselines'),(640,'03','Build an expert\nreview interface')]:
 rect(x,151,284,222,'white');tx(n,x+18,169,246,18,'teal',True);tx(t,x+18,224,246,27,bold=True)
banner('Deliver specimen images, locations, proposed labels, and uncertainty.')
end('Next steps and deliverable',20,'Next we will finalize categories, inspect crops, and freeze the evaluation protocol before training the baselines. The selected approach will support a simple interface linking specimen images and locations to proposed labels and uncertainty. The goal is useful scientific review, with clear evidence about where automated identification succeeds or fails.', 'Report p. 10, next steps and limitations. No promise of a finalized grouping, trained classifier, or completed viewer.')
for page,t,x,y,w,h in bounds:assert x>=0 and y>=0 and x+w<=W+1 and y+h<=H+1,(page,t,y,h)
prs.save(P/'Smithsonian_initial_presentation.pptx');pdf.save()
count=sum(len(re.findall(r"\b[\w’-]+\b",n['script'])) for n in notes)
(P/'speaker_notes.md').write_text('# Under-five-minute speaker guide\n\nTen slides. Planned delivery: 4:20 (260 seconds). Spoken script: '+str(count)+' words. At 125 words/minute this is approximately '+str(round(count/125,1))+' minutes before pauses. Rehearse once; timing is a target, not a guarantee. Source notes are not spoken.\n\n'+'\n\n'.join(f"## Slide {n['slide']} — {n['title']} ({n['seconds']} seconds)\n\n{n['script']}\n\nPresenter reference: {n['source']}" for n in notes))
(P/'timing.json').write_text(json.dumps({'slides':10,'target_seconds':sum(n['seconds'] for n in notes),'spoken_words':count,'slides_timing':notes},indent=2))
doc=fitz.open(P/'Smithsonian_initial_presentation.pdf');thumbs=[]
for i,page in enumerate(doc):
 path=P/f'slide_{i+1:02}.png';page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(path);thumbs.append(Image.open(path).resize((640,360)))
sheet=Image.new('RGB',(1280,1800),'white')
for i,im in enumerate(thumbs):sheet.paste(im,((i%2)*640,(i//2)*360))
sheet.save(P/'slide_overview.png')
print('Slides:',len(prs.slides),'Spoken words:',count,'Target seconds:',sum(n['seconds'] for n in notes))
