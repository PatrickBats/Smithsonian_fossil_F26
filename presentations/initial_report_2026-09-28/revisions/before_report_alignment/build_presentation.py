from pathlib import Path
import textwrap,json,hashlib,re
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image
import fitz
P=Path(__file__).resolve().parent;R=P.parents[1];A=R/'reports/initial_report_complete_2026-09-27/figures/specimen_sources'
W,H=960,540
C={'bg':'F8F5EF','ink':'202D48','teal':'93451F','light':'EEE3D6','muted':'596477','white':'FFFFFF','gold':'B97832'}
for n,f in [('Lato','Lato-Regular.ttf'),('Lato-Bold','Lato-Bold.ttf')]:pdfmetrics.registerFont(TTFont(n,'/usr/share/fonts/truetype/lato/'+f))
prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
pdf=canvas.Canvas(str(P/'Smithsonian_initial_presentation.pdf'),pagesize=(W,H));notes=[];bounds=[]
def rect(x,y,w,h,c):
 c=C.get(c,c);sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();pdf.setFillColor('#'+c);pdf.rect(x,H-y-h,w,h,stroke=0,fill=1)
def tx(t,x,y,w,size=20,c='ink',bold=False):
 font='Lato-Bold' if bold else 'Lato';lines=[]
 for para in t.split('\n'):
  line=''
  for word in para.split():
   cand=(line+' '+word).strip()
   if line and pdfmetrics.stringWidth(cand,font,size)>w:lines.append(line);line=word
   else:line=cand
  lines.append(line)
 height=len(lines)*size*1.2+4;bounds.append((len(prs.slides),t[:40],x,y,w,height))
 sh=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(height));tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Lato';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(c,c));p.line_spacing=Pt(size*1.2);p.space_after=Pt(0)
 pdf.setFont(font,size);pdf.setFillColor('#'+C.get(c,c))
 for i,line in enumerate(lines):pdf.drawString(x,H-y-size-i*size*1.2,line)
def pic(path,x,y,w,h):
 iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale;x+=(w-ww)/2;y+=(h-hh)/2
 s.shapes.add_picture(str(path),Pt(x),Pt(y),width=Pt(ww),height=Pt(hh));pdf.drawImage(str(path),x,H-y-hh,ww,hh)
def new(k,title):
 global s
 s=prs.slides.add_slide(prs.slide_layouts[6]);rect(0,0,W,H,'bg');rect(36,28,25,4,'teal');tx(k,73,20,850,11,'teal',True);tx(title,36,57,888,32,bold=True)
 rect(36,505,888,1,'light');tx('SMITHSONIAN × RICE D2K  /  INITIAL PROJECT PRESENTATION',36,516,830,9,'muted');tx(str(len(prs.slides)),900,514,24,11,'teal',True)
def end(title,seconds,spoken,source):
 n=f'TARGET: {seconds} seconds\n\nSAY:\n{spoken}\n\nSOURCE / PRESENTER CONTEXT (not spoken):\n{source}'
 s.notes_slide.notes_text_frame.text=n;notes.append(dict(slide=len(prs.slides),title=title,seconds=seconds,script=spoken,source=source));pdf.showPage()
def banner(t):rect(36,444,888,43,'light');tx(t,50,454,860,18,'teal',True)
new('01 / THE PROJECT','Identifying fossil pollen at scale')
tx('Small fossils.\nA large scientific record.',36,157,435,36,bold=True)
pic(A/'shard_0_target_027_crop.png',559,129,324,290)
tx('Smithsonian × Rice D2K',36,301,450,23,'teal',True)
tx('Patrick Batsell • Yun Ying Tsai • Yeonju Kim\nYunfan Bao • Alan Yang',36,352,478,16,'muted')
end('Opening',20,'Our project uses machine learning to help identify fossil pollen in Smithsonian microscope images. These tiny fossils form a large scientific record, but examining them requires substantial expert time. We are developing a workflow that connects specimen detection with category identification.', 'Uploaded report, Introduction. Real focus-stacked Bombacacidites crop; expert annotation category.')
new('02 / WHY IT MATTERS','Pollen reveals ancient vegetation')
tx('~50 million\nyears ago',36,158,443,47,'teal',True)
tx('Identify past plant communities.\nSupport research on warm climates.',36,313,460,24)
pic(A/'shard_0_target_008_crop.png',589,137,285,285)
end('Scientific motivation',20,'Pollen can survive in sediments long after the plants that produced it disappear. Identifying those fossils helps researchers reconstruct vegetation. Here, the scientific motivation is understanding North American plant communities around fifty million years ago and their relationship to past warm climates.', 'Uploaded report, Introduction. Image: expert-labeled Platycarya platycarioides; not a model prediction.')
new('03 / THE TASK','Finding a specimen is not the same as naming it')
for x,head,body in [(36,'DETECTION','Where is the specimen?'),(504,'CLASSIFICATION','What category is it?')]:
 rect(x,170,420,184,'white');tx(head,x+24,192,372,16,'teal',True);tx(body,x+24,247,365,30,bold=True)
banner('We will compare two ways of connecting these tasks.')
end('Problem definition',20,'There are two distinct tasks. Detection locates a specimen in an image. Classification identifies its category from its appearance. Previous project work provides a detection foundation. Our main question is whether identification works better inside the detector or in a separate classifier.', 'Uploaded report, Objectives and Proposed Models.')
new('04 / THE DATA','83 slides. 2,290 target annotations.')
for x,num,label in [(36,'83','annotated slides'),(338,'2,290','target annotations'),(640,'32','candidate categories')]:
 rect(x,168,284,162,'white');tx(num,x+20,189,245,48,'teal',True);tx(label,x+20,273,245,21,'muted')
tx('Large images with multiple focal depths',36,372,888,26,bold=True)
banner('Images stay on Rice’s cluster; we read small regions.')
end('Dataset',25,'The working dataset contains eighty-three annotated slides. Our initial target inventory has two thousand two hundred ninety annotations across thirty-two candidate categories, within a larger set of eleven thousand two hundred seventy-five named annotations. The final category grouping remains to be fixed. Images contain multiple focal depths, so we read small regions on Rice’s cluster.', 'Uploaded report, Data. Counts precede quality review, duplicate checks, category regrouping and sampling.')
new('05 / THE CHALLENGE','Some categories have far fewer examples')
for y,label,n in [(150,'Bisaccate',316),(277,'Cupuliferoipollenites sp.',27)]:
 tx(label,36,y,740,25,bold=True);rect(36,y+48,754,29,'light');rect(36,y+48,754*n/316,29,'teal');tx(str(n),815,y+36,100,37,'teal',True)
banner('Evaluate each category—not just overall accuracy.')
end('Imbalance',25,'Examples are unevenly distributed. Bisaccate has three hundred sixteen annotations, while the smallest target category has twenty-seven. Slide coverage matters too: many examples from one slide do not provide the same diversity as examples across several slides. We will consider training-only balancing strategies and evaluate performance separately for each category.', 'Uploaded report, Class balance. Bars use a common linear scale. Candidate strategies: undersampling, balanced sampling and weighted loss.')
new('06 / APPROACH A','One model finds and identifies')
for x,w,t,col in [(36,213,'Image\ntiles','light'),(296,352,'Multiclass RF-DETR\nLocation + category','teal'),(695,229,'Labeled\nspecimens','light')]:
 rect(x,191,w,132,col);tx(t,x+20,221,w-40,26,'white' if col=='teal' else 'ink',True)
for x in [254,653]:tx('→',x,233,38,31,'teal',True)
banner('A unified model learns both tasks together.')
end('Joint approach',30,'The first approach uses multiclass RF-DETR, a transformer-based detector. It receives image tiles and predicts both specimen locations and categories. Predictions from overlapping tiles are then merged in slide coordinates. The attraction is a unified model, but learning rare or visually similar categories alongside localization may be difficult. This is a proposed experiment.', 'Uploaded report, Model 1. Simplified diagram; focus stacking, coordinate mapping and duplicate merging remain part of the pipeline.')
new('07 / APPROACH B','Find first. Classify the crop.')
for x,w,t,col in [(36,250,'RF-DETR\nFind specimens','teal'),(341,225,'Specimen\ncrops','light'),(621,303,'Swin-Tiny\nIdentify the category','teal')]:
 rect(x,191,w,132,col);tx(t,x+19,220,w-38,25,'white' if col=='teal' else 'ink',True)
for x in [297,577]:tx('→',x,234,35,29,'teal',True)
banner('Separates identification quality from detection quality.')
end('Two-stage approach',30,'The second approach uses RF-DETR to find specimens, then passes individual crops to a Swin-Tiny image classifier. This lets us examine classification separately. However, missed specimens and poor crops can affect the final result. We will test the classifier on both expert-defined crops and detector-generated crops to identify where errors arise. This approach is also still proposed.', 'Uploaded report, Model 2. Same input preprocessing and split as Model 1; detector boxes mapped and merged before crop classification.')
new('08 / PRELIMINARY EVIDENCE','The existing detector can recover known targets')
for x,n,lab in [(36,'54 / 64','Best focal plane'),(489,'51 / 64','Focus stack')]:
 rect(x,157,435,181,'white');tx(n,x+23,177,390,57,'teal',True);tx(lab,x+23,263,390,25,bold=True)
tx('YOLO checkpoints • 64 targets across 29 slides',36,368,888,22,'muted')
banner('A feasibility check—not independent classification performance.')
end('Pilot evidence',30,'We tested two existing YOLO detector checkpoints on sixty-four known targets from twenty-nine slides. The best-plane model recovered fifty-four targets, and the focus-stacked model recovered fifty-one under the same confidence and box-overlap rule. These results demonstrate feasibility. They are not classification accuracy or independent generalization: crops were centered on known targets, and many slides appear in historical training records.', 'Uploaded report, Preliminary Detector Assessment. Confidence >=0.50, IoU >=0.50. Different checkpoint weights; cannot isolate preprocessing. No RF-DETR/Swin classification training reported.')
new('09 / EVALUATION','A fair comparison needs unseen source groups')
for i,(h,b) in enumerate([('Separate slides','Keep each specimen and all its views together.'),('Use one protocol','Same categories, preprocessing, and evaluation data.'),('Count all errors','Measure missed specimens and wrong categories.')]):
 y=146+i*89;tx(str(i+1).zfill(2),36,y,60,28,'teal',True);tx(h,112,y,285,25,bold=True);tx(b,425,y+2,490,22,'muted')
end('Fair testing',30,'Both approaches will use the same category definitions, preprocessing, and evaluation groups. We will separate slides and investigate related samples to reduce leakage. Category-level metrics will reveal weaknesses hidden by overall accuracy. End-to-end evaluation must count missed specimens as well as wrong labels. Measuring detection precision also requires reviewed regions with sufficiently complete annotations.', 'Uploaded report, Partitioning and Evaluation. Metrics include macro-F1, balanced accuracy, per-class recall, and end-to-end detection measures where labels support them.')
new('10 / NEXT STEPS','From annotated specimens to useful predictions')
for x,n,t in [(36,'01','Inspect crops\nand define groups'),(338,'02','Freeze the split\nand compare models'),(640,'03','Build a simple\nreview interface')]:
 rect(x,159,284,211,'white');tx(n,x+19,179,245,18,'teal',True);tx(t,x+19,235,245,27,bold=True)
banner('Deliver images, proposed categories, and uncertainty for expert review.')
end('Close',25,'Next, we will inspect crops, define the category groups, and freeze the evaluation split. Then we can compare the two baselines and use their errors to guide development. Our intended deliverable is a simple interface linking each specimen image and location to proposed labels and uncertainty, so researchers can review results efficiently.', 'Uploaded report, Next Steps and Objectives. No completed viewer or trained classifier is claimed.')
for page,t,x,y,w,h in bounds:assert x>=0 and y>=0 and x+w<=W+1 and y+h<=H+1,(page,t,y,h)
prs.save(P/'Smithsonian_initial_presentation.pptx');pdf.save()
count=sum(len(re.findall(r"\b[\w’-]+\b",n['script'])) for n in notes)
(P/'speaker_notes.md').write_text('# Under-five-minute speaker guide\n\nTen slides. Planned delivery: 4:15 (255 seconds). Spoken script: '+str(count)+' words. At 125 words/minute this is approximately '+str(round(count/125,1))+' minutes before pauses. Rehearse once; timing is a target, not a guarantee. Source notes are not spoken.\n\n'+'\n\n'.join(f"## Slide {n['slide']} — {n['title']} ({n['seconds']} seconds)\n\n{n['script']}\n\nPresenter reference: {n['source']}" for n in notes))
(P/'timing.json').write_text(json.dumps({'slides':10,'target_seconds':sum(n['seconds'] for n in notes),'spoken_words':count,'slides_timing':notes},indent=2))
doc=fitz.open(P/'Smithsonian_initial_presentation.pdf');thumbs=[]
for i,page in enumerate(doc):
 path=P/f'slide_{i+1:02}.png';page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(path);thumbs.append(Image.open(path).resize((640,360)))
sheet=Image.new('RGB',(1280,1800),'white')
for i,im in enumerate(thumbs):sheet.paste(im,((i%2)*640,(i//2)*360))
sheet.save(P/'slide_overview.png')
print('Slides:',len(prs.slides),'Spoken words:',count,'Target seconds:',sum(n['seconds'] for n in notes))
