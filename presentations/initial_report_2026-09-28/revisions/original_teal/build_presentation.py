from pathlib import Path
import textwrap,json,hashlib,re
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image
import fitz
P=Path(__file__).resolve().parent;R=P.parents[1];A=R/'reports/initial_report_complete_2026-09-27/figures/specimen_sources'
W,H=960,540
C={'bg':'F5F5F0','ink':'143842','teal':'087F83','light':'E5F0ED','muted':'576C73','white':'FFFFFF','gold':'B58237'}
for n,f in [('Lato','Lato-Regular.ttf'),('Lato-Bold','Lato-Bold.ttf')]:pdfmetrics.registerFont(TTFont(n,'/usr/share/fonts/truetype/lato/'+f))
prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
pdf=canvas.Canvas(str(P/'Smithsonian_initial_presentation.pdf'),pagesize=(W,H));notes=[];bounds=[]
def rect(x,y,w,h,c):
 c=C.get(c,c);sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();pdf.setFillColor('#'+c);pdf.rect(x,H-y-h,w,h,stroke=0,fill=1)
def tx(t,x,y,w,size=20,c='ink',bold=False):
 font='Lato-Bold' if bold else 'Lato';lines=[]
 for para in t.split('\n'):
  line=''
  for word in para.split():
   cand=(line+' '+word).strip()
   if line and pdfmetrics.stringWidth(cand,font,size)>w:lines.append(line);line=word
   else:line=cand
  lines.append(line)
 height=len(lines)*size*1.2+4;bounds.append((len(prs.slides),t[:40],x,y,w,height))
 sh=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(height));tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Lato';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(c,c));p.line_spacing=Pt(size*1.2);p.space_after=Pt(0)
 pdf.setFont(font,size);pdf.setFillColor('#'+C.get(c,c))
 for i,line in enumerate(lines):pdf.drawString(x,H-y-size-i*size*1.2,line)
def pic(path,x,y,w,h):
 iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale;x+=(w-ww)/2;y+=(h-hh)/2
 s.shapes.add_picture(str(path),Pt(x),Pt(y),width=Pt(ww),height=Pt(hh));pdf.drawImage(str(path),x,H-y-hh,ww,hh)
def new(k,title):
 global s
 s=prs.slides.add_slide(prs.slide_layouts[6]);rect(0,0,W,H,'bg');rect(36,28,25,4,'teal');tx(k,73,20,850,11,'teal',True);tx(title,36,57,888,32,bold=True)
 rect(36,505,888,1,'light');tx('SMITHSONIAN × RICE D2K  /  INITIAL PROJECT PRESENTATION',36,516,830,9,'muted');tx(str(len(prs.slides)),900,514,24,11,'teal',True)
def end(title,seconds,spoken,source):
 n=f'TARGET: {seconds} seconds\n\nSAY:\n{spoken}\n\nSOURCE / PRESENTER CONTEXT (not spoken):\n{source}'
 s.notes_slide.notes_text_frame.text=n;notes.append(dict(slide=len(prs.slides),title=title,seconds=seconds,script=spoken,source=source));pdf.showPage()
def banner(t):rect(36,444,888,43,'light');tx(t,50,454,860,18,'teal',True)
new('01 / THE QUESTION','From fossil pollen to past vegetation')
tx('Can we help identify\nmicroscopic fossils\nat collection scale?',36,139,420,31,bold=True)
tx('North American vegetation\n~50 million years ago',36,275,420,22,'muted')
for name,file,x in [('Arecipites','shard_0_target_016_crop.png',495),('Bombacacidites','shard_0_target_027_crop.png',711)]:
 pic(A/file,x,139,200,200);tx(name,x,351,210,17,'teal',True)
tx('Real specimens • expert annotation labels',495,386,425,13,'muted')
banner('Our task: find specimens, identify their category, support expert review.')
end('Why this matters',40,'Fossil pollen helps researchers understand which plants grew in ancient environments. Our Smithsonian project focuses on North American vegetation around fifty million years ago. The difficulty is scale: a microscope slide can contain many specimens mixed with debris, and identifying them requires expert attention. We want to support that work by connecting two tasks: finding a specimen and identifying its category. These are real examples from our dataset. The intended result is a tool that lets researchers inspect the image and review the proposed identification.', 'Uploaded report, Introduction and specimen figure. Images are focus-stacked crops, shown at different field widths; category names come from expert annotations.')
new('02 / THE DATA','Rich images. Uneven category coverage.')
for x,num,label in [(36,'83','annotated slides'),(338,'2,290','target annotations'),(640,'32','candidate target categories')]:
 rect(x,132,284,115,'white');tx(num,x+18,141,250,42,'teal',True);tx(label,x+18,204,250,17,'muted')
tx('Multiple focal depths per slide',36,276,430,23,bold=True)
tx('Read small image regions on Rice’s cluster; preserve the original labels and locations.',36,316,420,20,'muted')
for y,label,n in [(278,'Bisaccate',316),(339,'Cupuliferoipollenites sp.',27)]:
 tx(label,508,y,355,17,bold=True);rect(508,y+28,345,12,'light');rect(508,y+28,345*n/316,12,'teal');tx(str(n),868,y+17,60,20,'teal',True)
banner('11,275 named annotations overall; 2,290 are in the initial target groups.')
end('Data and challenge',45,'We have eighty-three annotated slides. The complete inventory contains eleven thousand two hundred seventy-five named annotations, and the initial classification task draws on two thousand two hundred ninety annotations across thirty-two candidate categories. The final grouping still needs to be fixed. Each slide contains multiple focal depths, so we process small regions on the cluster rather than loading an entire slide into memory. Coverage is uneven: Bisaccate has three hundred sixteen annotations, while the smallest target category has twenty-seven. We also need to consider how many independent slides each category spans.', 'Uploaded report, Data and Initial Exploration; verified category_summary.csv. Counts precede quality review and sampling. Bars share a common scale; the two displayed categories illustrate the range, not the full distribution.')
new('03 / THE COMPARISON','One joint model—or two specialized stages?')
tx('A / JOINT DETECTION + IDENTIFICATION',36,129,870,14,'teal',True)
for x,w,t,col in [(36,190,'Focus-stacked\nimage tiles','light'),(273,307,'Multiclass RF-DETR\nFind + identify','teal'),(627,297,'Locations + categories\n+ scores','light')]:
 rect(x,157,w,87,col);tx(t,x+14,178,w-28,21,'white' if col=='teal' else 'ink',True)
tx('→',235,180,33,27,'teal',True);tx('→',590,180,33,27,'teal',True)
tx('B / DETECTION FOLLOWED BY CLASSIFICATION',36,278,870,14,'teal',True)
for x,w,t,col in [(36,154,'Image\ntiles','light'),(221,192,'RF-DETR\nFind specimens','teal'),(444,125,'Specimen\ncrops','light'),(600,187,'Swin-Tiny\nIdentify type','teal'),(818,106,'Labeled\nresults','light')]:
 rect(x,309,w,87,col);tx(t,x+10,329,w-20,18,'white' if col=='teal' else 'ink',True)
for x in [192,415,571,789]:tx('→',x,335,27,22,'teal',True)
banner('Same categories, slide split, preprocessing, and evaluation rules.')
end('Two proposed architectures',60,'Our central question is whether finding and identifying specimens should be learned together or separately. In approach A, multiclass RF-DETR predicts both a specimen’s location and its category. In approach B, RF-DETR first locates specimens, and a separate Swin-Tiny image classifier examines each crop to identify its type. The first approach offers a unified model. The second lets us study classification independently, but detection errors can carry into the classifier. Both are transformer-based approaches. We will compare them using the same categories, slide split, and preprocessing. We will also test the separate classifier on expert-defined crops to distinguish classification problems from detection problems.', 'Uploaded report, Proposed Models. Simplified diagrams omit mapping and duplicate merging; both pipelines require them. Neither proposed classification approach has been trained in the reported work. RF-DETR = Roboflow Detection Transformer; Swin is a hierarchical shifted-window vision transformer.')
new('04 / STARTING EVIDENCE','Existing detection models run on our images')
tx('64 known targets • 29 slides • two frozen YOLO checkpoints',36,119,888,20,'muted')
for x,n,lab in [(36,'54 / 64','Best focal plane'),(489,'51 / 64','Focus stack')]:
 rect(x,184,435,166,'white');tx(n,x+23,202,390,56,'teal',True);tx(lab,x+23,280,390,23,bold=True)
tx('Target recovery at confidence ≥ 0.50 and box overlap (IoU) ≥ 0.50',36,371,888,17,'muted')
banner('Feasibility evidence—not classification accuracy or an independent test.')
end('What is completed',40,'We have already run a small diagnostic with two existing YOLO detector checkpoints. Across sixty-four known targets from twenty-nine slides, the best-plane model recovered fifty-four targets and the focus-stacked model recovered fifty-one under the same matching rule. This establishes that the existing detection software works on sampled regions from our images. It is not classification accuracy or an independent test: the crops were centered on known specimens, and many slides appear in historical training records. The RF-DETR and Swin comparison remains proposed work.', 'Uploaded report, Preliminary Detector Assessment. Recovery requires confidence >=0.5 and IoU >=0.5. Historical membership: 50 train, 9 test, 5 unresolved; checkpoint exposure not fully verified. Different weights mean the result does not isolate preprocessing.')
new('05 / A FAIR TEST','Measure biological recognition—not slide familiarity')
for x,num,head,body in [(36,'01','Separate source slides','Keep each specimen and all its focal views in one data partition.'),(338,'02','Evaluate every category','Use macro-F1, balanced accuracy, and per-category error analysis.'),(640,'03','Test the full workflow','Count missed specimens as well as incorrect identifications.')]:
 rect(x,146,284,259,'white');tx(num,x+18,164,245,17,'teal',True);tx(head,x+18,203,245,24,bold=True);tx(body,x+18,279,245,19,'muted')
banner('Training-only balancing; fixed validation and test sets.')
end('Evaluation',45,'A reliable comparison must test biological recognition rather than familiarity with a slide’s appearance. We will keep specimens and all their views together, separate source slides, and investigate related samples that may need additional grouping. We will measure performance for each category instead of relying on overall accuracy. Candidate imbalance strategies include undersampling abundant categories, balanced sampling, and weighted losses, applied only to training. Finally, we will evaluate the full workflow, including missed specimens. Detection precision requires reviewed regions with complete labels; an unmarked object is not automatically a false detection.', 'Uploaded report, Partitioning and Model Evaluation. Macro-F1 gives categories equal weight in averaging precision-recall summaries; balanced accuracy averages category recall. Related slides and checkpoint training exposure require review.')
new('06 / THE NEXT STEP','Build a useful, inspectable classification pipeline')
for i,(head,body) in enumerate([('Prepare','Inspect crops and agree on category groups.'),('Compare','Freeze the split and run both baselines.'),('Deliver','Show the specimen, proposed label, and uncertainty.')]):
 y=141+i*91;rect(36,y,48,48,'teal');tx(str(i+1),52,y+8,25,25,'white',True);tx(head,110,y-1,180,25,bold=True);tx(body,310,y+2,610,23,'muted')
banner('Goal: faster scientific review, with evidence about where models fail.')
end('Close',30,'Our next steps are to inspect the specimen crops, agree on category groups, and freeze the evaluation split. We will then run the two baselines and use their errors to decide which approach to develop further. The final output should connect a specimen’s image and location to a proposed category and uncertainty information. Success means giving scientists a useful way to review more material, while making the model’s limitations visible.', 'Uploaded report, Objectives and Next Steps. This is an initial project proposal, not a claim that a trained classifier or viewer is complete.')
for page,t,x,y,w,h in bounds:assert x>=0 and y>=0 and x+w<=W+1 and y+h<=H+1,(page,t,y,h)
prs.save(P/'Smithsonian_initial_presentation.pptx');pdf.save()
count=sum(len(re.findall(r"\b[\w’-]+\b",n['script'])) for n in notes)
(P/'speaker_notes.md').write_text('# Under-five-minute speaker guide\n\nSix slides. Planned delivery: 4:20 (260 seconds). Spoken script: '+str(count)+' words. At 125 words/minute this is approximately '+str(round(count/125,1))+' minutes before pauses. Rehearse once; timing is a target, not a guarantee. Source notes are not spoken.\n\n'+'\n\n'.join(f"## Slide {n['slide']} — {n['title']} ({n['seconds']} seconds)\n\n{n['script']}\n\nPresenter reference: {n['source']}" for n in notes))
(P/'timing.json').write_text(json.dumps({'slides':6,'target_seconds':sum(n['seconds'] for n in notes),'spoken_words':count,'slides_timing':notes},indent=2))
doc=fitz.open(P/'Smithsonian_initial_presentation.pdf');thumbs=[]
for i,page in enumerate(doc):
 path=P/f'slide_{i+1:02}.png';page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(path);thumbs.append(Image.open(path).resize((640,360)))
sheet=Image.new('RGB',(1280,1080),'white')
for i,im in enumerate(thumbs):sheet.paste(im,((i%2)*640,(i//2)*360))
sheet.save(P/'slide_overview.png')
print('Slides:',len(prs.slides),'Spoken words:',count,'Target seconds:',sum(n['seconds'] for n in notes))
