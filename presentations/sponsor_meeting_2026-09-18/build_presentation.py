from pathlib import Path
import csv, textwrap, json, hashlib
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image
P=Path(__file__).parent; ROOT=P.parents[1]
W,H=960,540
C={'bg':'F6F4EE','ink':'142F38','muted':'596D71','teal':'16756D','light':'E3EEEA','gold':'B58132','sand':'F2E5CB','white':'FFFFFF','rule':'D6DDD8'}
for name,file in [('Lato','Lato-Regular.ttf'),('Lato-Bold','Lato-Bold.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/lato/'+file))
prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
pdf=canvas.Canvas(str(P/'Smithsonian_sponsor_meeting.pdf'),pagesize=(W,H));pdf.setTitle('Smithsonian fossil pollen — sponsor discussion')
notes=[]; texts=[]; bounds=[]
def rgb(h):return RGBColor.from_string(h)
def rect(x,y,w,h,color):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(C.get(color,color));sh.line.fill.background()
 pdf.setFillColor('#'+C.get(color,color));pdf.rect(x,H-y-h,w,h,stroke=0,fill=1)
def tx(text,x,y,w,size=20,color='ink',bold=False,leading=1.22):
 font='Lato-Bold' if bold else 'Lato';lines=[]
 for para in text.split('\n'):
  line=''
  for word in para.split():
   cand=(line+' '+word).strip()
   if line and pdfmetrics.stringWidth(cand,font,size)>w-3:lines.append(line);line=word
   else:line=cand
  lines.append(line)
 height=len(lines)*size*leading+6
 box=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(height));tf=box.text_frame;tf.clear();tf.word_wrap=False
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Lato';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(C.get(color,color));p.space_after=Pt(0);p.line_spacing=Pt(size*leading)
 pdf.setFont(font,size);pdf.setFillColor('#'+C.get(color,color))
 for i,line in enumerate(lines):pdf.drawString(x,H-y-size-i*size*leading,line)
 bounds.append((len(prs.slides),text[:35],x,y,w,height));texts.append(text)
 return height

def pic(path,x,y,w,h):
 iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale;x+=(w-ww)/2;y+=(h-hh)/2
 s.shapes.add_picture(str(path),Pt(x),Pt(y),width=Pt(ww),height=Pt(hh));pdf.drawImage(str(path),x,H-y-hh,ww,hh)
def new(kicker,title,subtitle=''):
 global s
 s=prs.slides.add_slide(prs.slide_layouts[6]);rect(0,0,W,H,'bg');rect(40,33,26,4,'teal');tx(kicker.upper(),77,26,820,11,'teal',True)
 tx(title,40,63,880,31,bold=True)
 if subtitle:tx(subtitle,40,111,875,16,'muted')
 rect(40,505,880,1,'rule');tx('SMITHSONIAN × RICE D2K   /   SPONSOR DISCUSSION   /   18 SEPT 2026',40,515,800,9,'muted');tx(f'{len(prs.slides):02}',891,512,30,12,'teal',True)
def end(title,note):
 s.notes_slide.notes_text_frame.text=note;notes.append((title,note));pdf.showPage()
def question(text,y=449):
 rect(40,y,880,43,'light');tx(text,55,y+10,850,17,'teal',True)
def card(x,y,w,h,num,title,body):
 rect(x,y,w,h,'white');tx(num,x+18,y+14,w-36,14,'teal',True);tx(title,x+18,y+43,w-36,23,bold=True);tx(body,x+18,y+84,w-36,17,'muted')
img=P/'assets/alaska_boxes.png'
new('01 / Purpose','From fossil pollen to a picture of past vegetation','Project progress and decisions for our first classifier')
pic(img,530,153,390,337)
tx('Identify pollen types\nin microscope images.',40,179,455,34,bold=True)
tx('Support research on North American vegetation about 50–60 million years ago.',40,282,430,22,'muted')
tx('Today: what we have → what we need to agree',40,417,440,17,'teal',True)
end('Opening — 45 seconds',"OPEN: Thank you for meeting with us. Our understanding is that the project will help identify fossil pollen types in microscope images, supporting your work on past vegetation and warm climates. Today we will show what we can access, explain how we have organized the labels, and agree on the first useful version. We have not trained a classifier yet. The image is a real region from a slide on Rice’s cluster, with existing annotation boxes—not model predictions.\nMEETING FLOW: Allow about eight minutes for the overview and use the remaining time for discussion and decisions.\nSOURCE: Sponsor presentation and September 11 meeting; Alaska preview provenance.")
new('02 / Shared understanding','Finding a grain and naming it are different tasks')
card(40,170,265,219,'PREVIOUS WORK','Find the specimen','Locate an object within a large microscope image.')
card(347,170,265,219,'THIS SEMESTER','Identify its type','Assign a useful category to an already located specimen.')
card(654,170,266,219,'SCIENTIFIC USE','Review the result','Let an expert inspect the image, category and uncertain cases.')
tx('→',315,249,25,25,'teal',True);tx('→',622,249,25,25,'teal',True)
question('Does this capture the most useful focus for this semester?')
end('Agree on the problem — 45 seconds',"SAY: The previous group focused on where palynomorphs are. Our main focus is what type each specimen is. Those are separate tasks. A box can show where to look without telling us a detailed biological category. We propose beginning with existing marked specimens, then connecting the classifier to detection results.\nASK: Does this reflect the priority for this semester?\nLISTEN FOR: Whether sponsors expect classification of expert annotations first, existing detector outputs first, or a complete whole-slide system. Do not promise a complete new detector or a full slide application.\nTRANSITION: Here is what those marked specimens look like in our actual data.\nSOURCE: September 11 sponsor meeting and project context.")
new('03 / A real example','We can view small regions of the large slides')
pic(img,40,137,440,354)
tx('The boxes tell us where to look.',525,167,382,26,bold=True)
tx('These examples are labeled “pol” (broad pollen), not a detailed type.',525,247,372,22,'muted')
rect(510,339,410,109,'light');tx('Images stay on Rice’s cluster.',526,352,375,20,'teal',True);tx('We read small regions instead of downloading whole slides.',526,384,375,18)
end('Explain the image — 60 seconds',"POINT: The blue boxes surround four annotated objects in a small image region. This entire panel is only a tiny part of one slide. The full Alaska image is about 28.5 GB. We can read a small region without bringing the whole slide onto our laptops. The original data stay on Rice’s computing cluster, NOTS.\nIMPORTANT: These four objects have the broad original title pol. They illustrate object location, not four confirmed examples of our detailed target categories. The displayed image is an existing focus-stacked tile from Spring 2026; different focus levels can help reveal structure. We have not established the best focus representation for classification.\nSOURCE: C_418058_W_Nassichuk_R_2025_01_21_14_59_16_Alaska; HDF5 tile_63624_42509. Preview boxes were checked against annotation geometry; see assets/provenance.json.")
new('04 / Progress','The label inventory is ready','We connected slide files, specimen locations and names to our working categories.')
for x,n,label in [(40,'128','Annotation files linked to slide images'),(340,'933','Annotations in our YES working set'),(640,'27','Priority categories represented')]:
 rect(x,176,280,168,'white');tx(n,x+20,190,240,57,'teal',True);tx(label,x+20,269,240,19)
rect(40,365,880,71,'sand');tx('Next: extract and inspect the specimen images.',57,375,845,21,bold=True);tx('The labeled image dataset is not yet prepared for training.',57,405,845,16,'muted')
tx('Inventory completed • example images viewed • no classifier trained yet',40,460,870,18,'teal',True)
end('State progress accurately — 60 seconds',"SAY: We read 82 annotation files in the training folder and 46 additional annotation files, and found a matching image filename for each of the 128. This was an annotation and file inventory, not a visual inspection of all 128 slides. We mapped names to your category table, first using exact names, then formatting fixes, and then a set of provisional connections approved internally by our team. That leaves 933 annotation records in 27 of your YES categories. Of those 933, 112 depend on provisional connections.\nCAUTION: These are annotation records, not a claim that 933 unique specimens have passed a full quality review. The 27 categories are represented in the inventory; they are not all necessarily suitable for a first model. We have kept MAYBE categories separate.\nSOURCE: data/current/classification_annotations.csv and category_summary.csv, September 17 audit. The current working set supersedes earlier 819- and 821-record inventories.")
summary=list(csv.DictReader((ROOT/'data/archive/2026-09-17_provisional_working_snapshot/category_summary.csv').open()));by={r['category']:r for r in summary}
new('05 / What the inventory shows','Some categories have much stronger coverage','Examples from our current working inventory')
tx('CATEGORY',40,165,280,12,'muted',True);tx('ANNOTATIONS',361,165,180,12,'muted',True);tx('SLIDES',740,165,100,12,'muted',True)
for i,name in enumerate(['Bisaccate','Momipites sp','Caryapollenites sp.','Monocolpites sp.','Rhoipites sp.']):
 r=by[name];y=203+i*43;n=int(r['annotation_records']);tx(name,40,y-3,307,19,bold=True);rect(361,y,325,18,'light');rect(361,y,325*n/335,18,'teal');tx(str(n),695,y-5,45,18,'teal',True);tx(str(r['slides']),774,y-5,60,18,bold=True)
question('Which scarce categories are essential, and could more examples be labeled?')
end('Discuss a realistic first model — 90 seconds',"SAY: The examples are unevenly distributed. Bisaccate has 335 annotation records across 44 slides; Momipites has 262 across 40. Rhoipites has 39, but all are on one slide. A model can accidentally learn a slide’s appearance instead of the biological distinctions. Testing on different slides helps us assess whether it transfers. With only one slide for a category, we cannot put that category into separate training and test slides using this set.\nASK: Would a smaller first classifier covering well-supported categories be useful? Which categories matter most scientifically? Which scarce categories are essential, and could additional examples be labeled? Explain that few examples make learning variation and measuring performance harder.\nDO NOT: Present the displayed five as a finalized class selection or promise that count alone makes a category usable. Note that counts include provisional mappings.\nSOURCE: data/current/category_summary.csv, September 17 working selection.")
new('06 / Next work','Turn the inventory into a training dataset','We will use the agreed working categories and keep the original labels for traceability.')
card(40,171,280,223,'01','Extract the images','Read a small region around each selected annotation, including marks outside the old training areas.')
card(340,171,280,223,'02','Check the examples','Inspect position, focus and image quality; look for duplicate or unusable specimens.')
card(640,171,280,223,'03','Prepare a fair test','Choose supported categories and reserve slides for evaluation before training.')
question('Which visual features and preservation problems should we check first?')
end('Explain the next work — 90 seconds',"SAY: Our label inventory is ready, but we have not yet extracted and checked the complete set of specimen images. We will now extract regions around the selected annotations, including annotations outside the historical training rectangles. We will inspect the crops for alignment, focus, preservation, and duplicates, then prepare training and evaluation groups. The current count may change after this quality review.\nCATEGORY POLICY: We are using Ingrid’s category mappings and the agreed working extensions. Do not reopen the explicit Siltaria-to-Rhoipites or Liliacidites-to-Arecipites CSV mappings as if they were arbitrary. Original names and provisional flags remain recorded; these groupings are not claims of taxonomic synonymy. Ask about a specific qualified label only if it affects the work, rather than spending the main meeting reviewing every name.\nASK: Which visual features are essential for distinguishing the priority types? Which preservation or staining problems make a specimen unsuitable? Could you point out a few good and difficult examples?\nVERSION CHECK IF NEEDED: We can proceed with this working set. If there is a newer or corrected annotation set behind the spreadsheet, please identify it so we can record the version and compare it. Differences in totals have not been explained by the naming cleanup.\nSOURCE: Current working inventory and pipeline plan. Extraction, quality review and final evaluation splits remain future work.")
new('07 / Scientific priorities','What is most important for us to get right?')
card(40,171,280,223,'01','Critical distinctions','Which pollen types must we be especially careful not to confuse?')
card(340,171,280,223,'02','Fair testing','Check performance on specimens from slides the model did not learn from.')
card(640,171,280,223,'03','Uncertain cases','Allow a specimen to be flagged for expert review when the prediction is uncertain.')
question('For large categories, should ~60 specimens apply to training only?')
end('Agree on evaluation priorities — 3–4 minutes discussion',"SAY: We want to know whether the model learns useful biological distinctions, not just whether it recognizes familiar images. We propose holding some slides aside for evaluation, keeping all views of a specimen together, and reporting performance separately for each category. The exact split is not finalized.\nASK FIRST: What is of utmost importance for us to get correct in this project? Which distinctions are essential, and which mistakes would change your scientific conclusions? Is flagging an uncertain specimen preferable to forcing a label? Your email suggested about 60 specimens for categories over 100. Should that limit apply to the training examples, with separate specimens retained for evaluation?\nNOTE: Do not propose an arbitrary accuracy promise. We should agree on numeric targets after verifying labels and obtaining a first result. A confidence score is a model estimate, not automatically a calibrated probability.\nSOURCE: Sponsor email for the approximate-60 request; project pipeline plan for proposed evaluation approach.")
new('08 / Decision: deliverables','A simple result you can inspect')
card(40,175,280,201,'INPUT','A marked specimen','An image region with a known location on the original slide.')
card(340,175,280,201,'RESULT','A proposed category','The predicted type, with uncertain cases clearly flagged.')
card(640,175,280,201,'REVIEW','Image + identification','A simple viewer and a new annotation file; originals preserved.')
tx('First version: classify existing marked specimens.',40,404,880,22,bold=True)
question('Which viewing software and export format would fit your workflow?')
end('Confirm the handoff — 3–4 minutes discussion',"SAY: Our proposed first version takes an existing marked specimen, assigns a category, and lets you inspect the result beside the image. We would export predictions into a new annotation file and keep original expert annotations untouched. A simple viewer could show the category, location, and uncertain cases. This is a proposed workflow, not software we have already built.\nASK: Which application will you use to open the annotations? Can you supply an example of a result file that would be convenient? Should we prioritize annotated outputs, downloadable specimen crops, or browsing predictions? Is starting with existing marked specimens acceptable before integrating detector outputs?\nSCOPE: Full whole-slide navigation and interactive model fine-tuning remain possible extensions rather than promises for the first version.\nSOURCE: September 11 sponsor discussion and current pipeline plan.")
new('Appendix A / Inventory','All 27 represented YES categories','Current working annotation counts')
active=sorted([r for r in summary if r['priority']=='YES' and int(r['annotation_records'])>0],key=lambda r:-int(r['annotation_records']))
for col in range(2):
 x=40+col*450;tx('CATEGORY',x,158,315,11,'muted',True);tx('COUNT',x+318,158,53,11,'muted',True);tx('SLIDES',x+377,158,60,11,'muted',True)
 for j,r in enumerate(active[col*14:(col+1)*14]):
  y=183+j*20.5;tx(r['category'],x,y,315,13);tx(r['annotation_records'],x+325,y,45,13,'teal',True);tx(r['slides'],x+392,y,40,13)
end('Backup: full inventory',"Use this only if asked about a particular category. Counts reflect the current working mapping, including provisional aliases. A represented category is not necessarily ready for training or evaluation. Five YES categories have no matched records in this selection: Momipites flexus, Momipites ventifluminis, Nudopollis sp., Psilatricolpites sp., and Rousea sp. The 63 MAYBE annotations are separate. The 212 unmatched and 344 untitled entries are not part of the initial selection.\nSOURCE: data/current/category_summary.csv. Total 933 YES records across 27 represented categories.")
new('Final discussion','The main questions to leave with','Agree on priorities, next actions and who can help.')
items=[
 ('What matters most?','Which pollen types must we get right, and which mistakes matter most?'),
 ('Where do we need more examples?','Which scarce categories are essential? Could more specimens be labeled?'),
 ('What should we check in the images?','Which visual features and preservation problems should guide our checks?'),
 ('How should we start and test fairly?','Start with better-supported categories? Apply the ~60 limit to training only?'),
 ('What would be most useful to receive?','Which viewer, annotation format and review of uncertain cases fit your work?')]
for i,(head,body) in enumerate(items):
 y=151+i*58;rect(40,y,32,32,'teal');tx(str(i+1),50,y+4,22,18,'white',True);tx(head,88,y-2,824,19,bold=True);tx(body,88,y+23,824,16,'muted')
question('Before we finish: confirm an owner and next step for each open question.')
end('Final question recap — 2–3 minutes',"SAY: To close, these are the main questions we have discussed. What is most important for us to get right? Which rare categories are essential, and can additional examples be labeled? What should we look for when reviewing specimen images? Can we start with the better-supported categories, and should the approximately-60 limit apply to training only? Finally, what result format and way of reviewing uncertain cases would be useful?\nEMPHASIZE: The total of 933 annotations does not mean we have enough examples for every category. Scarce categories are harder to learn and evaluate, especially when all examples come from one slide. We should agree on a practical first scope and expand it as evidence allows.\nFACILITATE: Read back agreed answers rather than asking answered questions again. For remaining questions, record an owner, a next action and a follow-up date.\nNEXT: Extract and inspect the specimen crops, including annotations outside the old training rectangles; define evaluation groups; then train a first model.\nSOURCE: Recap of the discussion questions; no new sponsor commitments or model results.")
prs.save(P/'Smithsonian_sponsor_meeting.pptx');pdf.save()
(P/'speaker_notes.md').write_text('# Sponsor meeting speaker guide\n\nSeptember 18, 2026. Present slides 1–8, skip slide 9 unless the full inventory is needed, and close with slide 10. Aim for about 8 minutes of presentation plus discussion. Notes are also embedded in PowerPoint Presenter View.\n\n'+'\n\n'.join(f'## Slide {i+1}: {title}\n\n'+note.replace('\n','\n\n') for i,(title,note) in enumerate(notes)))
# Fail on any text block extending beyond the canvas.
for page,t,x,y,w,h in bounds:assert x>=0 and y>=0 and x+w<=W+1 and y+h<=H+1,(page,t,y,h)
assert len(prs.slides)==10
print('Created 10 slides, native editable text/shapes, speaker notes, PDF and Markdown guide.')
